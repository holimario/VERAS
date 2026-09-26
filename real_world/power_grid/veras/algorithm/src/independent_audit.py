#!/usr/bin/env python3
"""Independent NPZ replay; deliberately imports no construction-model code."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import openpyxl


SHAPES = {
    "uL": (63, 1), "uNS": (2, 2), "uEX": (2, 2), "zESS": (7, 1),
    "zPVG": (7, 1), "zWTG": (6, 1), "u_ess": (7, 96), "kt": (4, 96),
    "FF": (63, 4), "FS": (54, 4), "P": (63, 96), "Q": (63, 96),
    "l": (63, 96), "v": (54, 96), "P_s": (4, 96), "Q_s": (4, 96),
    "P_pv": (7, 96), "Q_pv": (7, 96), "P_wt": (6, 96), "Q_wt": (6, 96),
    "E_ess": (7, 96), "P_ess": (7, 96), "P_ess_cha": (7, 96),
    "P_ess_dis": (7, 96), "P_sop_1": (4, 96), "P_sop_2": (4, 96),
    "Q_sop_1": (4, 96), "Q_sop_2": (4, 96), "Pl_sop_1": (4, 96),
    "Pl_sop_2": (4, 96), "P_ls": (54, 96),
}
INTEGER = ("uL", "uNS", "uEX", "zESS", "zPVG", "zWTG", "u_ess", "kt")


def _mx(x) -> float:
    a = np.asarray(x, dtype=float)
    return float(np.max(a)) if a.size else 0.0


def _ma(x) -> float:
    a = np.asarray(x, dtype=float)
    return float(np.max(np.abs(a))) if a.size else 0.0


def _pos(x) -> float:
    return _mx(np.maximum(np.asarray(x, dtype=float), 0.0))


def audit(root: Path, solution_path: Path, tolerance: float = 5e-5) -> dict:
    spec = json.loads((root / "model_spec.json").read_text())
    wb = openpyxl.load_workbook(root / "line_load.xlsx", data_only=True, read_only=True)
    ws1, ws3 = wb["Sheet1"], wb["Sheet3"]
    branch = np.asarray([[ws1.cell(r, c).value for c in range(1, 5)] for r in range(2, 65)], float)
    bus = np.asarray([[ws1.cell(r, c).value for c in range(7, 9)] for r in range(2, 56)], float)
    fc = np.asarray([[ws3.cell(r, c).value for c in range(2, 5)] for r in range(2, 98)], float)
    wb.close()

    loaded = np.load(solution_path, allow_pickle=False)
    x: dict[str, np.ndarray] = {}
    structural_errors: list[str] = []
    for name, shape in SHAPES.items():
        if name not in loaded:
            structural_errors.append(f"missing array {name}")
            continue
        a = loaded[name]
        if a.shape != shape:
            structural_errors.append(f"{name} shape {a.shape}, expected {shape}")
        if a.dtype != np.float64:
            structural_errors.append(f"{name} dtype {a.dtype}, expected float64")
        if not np.all(np.isfinite(a)):
            structural_errors.append(f"{name} contains non-finite values")
        x[name] = np.asarray(a, dtype=float)
    if structural_errors:
        return {"passed": False, "tolerance": tolerance, "structural_errors": structural_errors}

    n_b, n_e, n_h = 54, 63, 96
    net, plan, op, econ, sets = spec["network"], spec["planning"], spec["operation"], spec["economics"], spec["sets"]
    fr, to, length = branch[:, 1].astype(int) - 1, branch[:, 2].astype(int) - 1, branch[:, 3]
    afr = np.zeros((n_b, n_e)); ato = np.zeros((n_b, n_e))
    afr[fr, np.arange(n_e)] = 1.0; ato[to, np.arange(n_e)] = 1.0
    incidence = ato - afr
    sbase = float(net["s_base_mva"]); vbase = float(net["v_base_kv"])
    zbase = vbase * vbase / sbase; ibase = sbase / vbase
    rr = float(net["resistance_ohm_per_km"]) * length / zbase
    xx = float(net["reactance_ohm_per_km"]) * length / zbase
    imax = float(net["line_current_limit_ka"]) / ibase
    isq = imax * imax; slimit = imax * float(net["v_max_pu"])
    load_qfac = np.tan(np.arccos(float(net["load_power_factor"])))
    dg_qfac = np.sqrt(1.0 - float(net["dg_min_power_factor"]) ** 2) / float(net["dg_min_power_factor"])
    pl = bus[:, 1, None] * fc[:, 2][None, :] / sbase; ql = pl * load_qfac
    sdg = float(plan["dg_unit_mw"]) / sbase
    ess_buses = np.asarray(sets["ess_buses"], int) - 1
    sess = np.full(7, float(op["ess_unit_mwh_non_idc"]) / sbase)
    sess[ess_buses == sets["idc_buses"][0] - 1] = float(op["ess_unit_mwh_idc"]) / sbase
    existing = np.asarray(sets["existing_line_rows"], int) - 1
    fault_e = np.asarray(sets["fault_line_rows"], int) - 1
    fault_h = np.asarray(spec["fault"]["fault_hours_one_based"], int) - 1
    mask = np.ones((n_e, n_h)); mask[np.ix_(fault_e, fault_h)] = 0.0
    shed_mask = np.zeros((n_b, n_h)); shed_mask[:, fault_h] = 1.0

    families: dict[str, dict] = {}

    def rec(name: str, residual: float, scale: float = 1.0) -> None:
        raw = max(0.0, float(residual))
        denom = max(1.0, float(scale))
        families[name] = {"max_abs_or_positive_part": raw, "scale": denom, "normalized": raw / denom}

    # Domains and all eight integral families.
    integ_by = {name: _ma(x[name] - np.rint(x[name])) for name in INTEGER}
    max_integrality = max(integ_by.values())
    rec("binary_and_integer_integrality", max_integrality)
    rec("binary_domains", max(_pos(-x["uL"]), _pos(x["uL"] - 1), _pos(-x["uNS"]), _pos(x["uNS"] - 1), _pos(-x["uEX"]), _pos(x["uEX"] - 1), _pos(-x["u_ess"]), _pos(x["u_ess"] - 1)))
    rec("integer_count_domains", max(_pos(1 - x["zESS"]), _pos(x["zESS"] - 20), _pos(1 - x["zPVG"]), _pos(x["zPVG"] - 50), _pos(1 - x["zWTG"]), _pos(x["zWTG"] - 50), _pos(-8 - x["kt"]), _pos(x["kt"] - 8)))
    nonnegative_names = ("FS", "l", "P_s", "P_pv", "P_wt", "E_ess", "P_ess_cha", "P_ess_dis", "Q_sop_1", "Q_sop_2", "Pl_sop_1", "Pl_sop_2", "P_ls")
    rec("continuous_nonnegative_domains", max(_pos(-x[n]) for n in nonnegative_names))

    # Planning.
    rec("existing_line_fixing", _ma(x["uL"][existing, 0] - 1.0))
    new_count = float(np.sum(x["uNS"]))
    rec("line_count_upper", _pos(np.sum(x["uL"]) + new_count - 52.0), 52.0)
    rec("topology_cardinality", _ma(np.repeat(np.sum(x["uL"]) + new_count - 52.0, 4)), 52.0)
    rec("ess_aggregate_limit", _pos(np.sum(x["zESS"]) - float(plan["ess_units_total_max"])), float(plan["ess_units_total_max"]))
    rec("pv_aggregate_limit", _pos(np.sum(x["zPVG"]) - float(plan["pv_units_total_max"])), float(plan["pv_units_total_max"]))
    rec("station_option_selection", max(_pos(np.sum(x["uEX"], axis=1) - 1), _pos(np.sum(x["uNS"], axis=1) - 1)))

    # Connectivity and exact expansion-gated roots.
    ul = x["uL"][:, 0]
    rec("FF_line_upper_lower", max(_pos(x["FF"] - 54.0 * ul[:, None]), _pos(-x["FF"] - 54.0 * ul[:, None])), 54.0)
    rec("FF_FS_balance", _ma(incidence @ x["FF"] - 1.0 + x["FS"]), 54.0)
    rec("FS_nonroot_zero", _ma(x["FS"][:50]))
    gates = np.r_[np.sum(x["uEX"], axis=1), np.sum(x["uNS"], axis=1)]
    rec("FS_expansion_gated_roots", max(_pos(x["FS"][50:] - 54.0 * gates[:, None]), _pos(-x["FS"][50:])), 54.0)

    # Stations and voltage/taps.
    scap = np.empty(4)
    scap[0] = 16.7 / sbase + np.dot(np.asarray(plan["existing_station_expansion_mva"][0]) / sbase, x["uEX"][0])
    scap[1] = 16.7 / sbase + np.dot(np.asarray(plan["existing_station_expansion_mva"][1]) / sbase, x["uEX"][1])
    scap[2] = np.dot(np.asarray(plan["candidate_station_sizes_mva"][0]) / sbase, x["uNS"][0])
    scap[3] = np.dot(np.asarray(plan["candidate_station_sizes_mva"][1]) / sbase, x["uNS"][1])
    derate = np.ones((4, 96)); derate[:2, fault_h] = float(spec["fault"]["substation_remaining_capacity_fraction"])
    rec("substation_capacity_cones_and_derating", _pos(np.hypot(x["P_s"], x["Q_s"]) - 0.5 * scap[:, None] * derate), _mx(scap))
    vmin2, vmax2, bigm = float(net["v_min_pu"]) ** 2, float(net["v_max_pu"]) ** 2, float(net["big_m"])
    rec("ordinary_bus_voltage_bounds", max(_pos(vmin2 - x["v"][:50]), _pos(x["v"][:50] - vmax2)), vmax2)
    rec("existing_station_voltage_tap_equalities", _ma(x["v"][50:52] - (1.01 + 0.025 * x["kt"][:2])), 1.21)
    new_v_res = 0.0
    for i in range(2):
        on = gates[2 + i]; vv = x["v"][52 + i]; target = 1.01 + 0.025 * x["kt"][2 + i]
        new_v_res = max(new_v_res, _pos(target - bigm * (1 - on) - vv), _pos(vv - target - bigm * (1 - on)), _pos(vmin2 - bigm * on - vv), _pos(vv - vmax2 - bigm * on), _pos(-8 * on - x["kt"][2 + i]), _pos(x["kt"][2 + i] - 8 * on))
    rec("candidate_station_voltage_and_tap_gates", new_v_res, bigm)

    # Branch mask/bounds/drop/SOC.
    yeff = ul[:, None] * mask
    rec("line_l_on_off_bounds", max(_pos(-x["l"]), _pos(x["l"] - isq * yeff)), isq)
    rec("line_PQ_on_off_bounds", max(_pos(np.abs(x["P"]) - slimit * yeff), _pos(np.abs(x["Q"]) - slimit * yeff)), slimit)
    drop = x["v"][fr] - x["v"][to] - 2 * rr[:, None] * x["P"] - 2 * xx[:, None] * x["Q"] + (rr * rr + xx * xx)[:, None] * x["l"]
    rec("voltage_drop_bigM", _pos(np.abs(drop) - bigm * (1 - yeff)), bigm)
    branch_soc = np.sqrt((2 * x["P"]) ** 2 + (2 * x["Q"]) ** 2 + (x["v"][fr] - x["l"]) ** 2) - (x["v"][fr] + x["l"])
    rec("branch_SOC_cones", _pos(branch_soc), _mx(x["v"][fr] + x["l"]))
    rec("fault_mask_scope", max(_ma(x["P"][np.ix_(fault_e, fault_h)]), _ma(x["Q"][np.ix_(fault_e, fault_h)]), _ma(x["l"][np.ix_(fault_e, fault_h)])))

    # Renewables.
    renewable_res = {}
    for tag, zname, pf, pgname, qgname in (("pv", "zPVG", fc[:, 1], "P_pv", "Q_pv"), ("wind", "zWTG", fc[:, 0], "P_wt", "Q_wt")):
        z = x[zname][:, 0]; pg, qg = x[pgname], x[qgname]
        avail = z[:, None] * pf[None, :] * sdg
        renewable_res[f"{tag}_availability"] = max(_pos(float(op["dg_min_dispatch_fraction"]) * avail - pg), _pos(pg - avail))
        renewable_res[f"{tag}_installed_capacity_cones"] = _pos(np.hypot(pg, qg) - z[:, None] * sdg)
        renewable_res[f"{tag}_actual_P_power_factor"] = _pos(np.abs(qg) - dg_qfac * pg)
    for name, val in renewable_res.items(): rec(name, val, 1.0)

    # Storage.
    cap = x["zESS"][:, 0, None] * sess[:, None]; pmax = float(op["ess_power_to_energy_ratio"]) * cap
    rec("storage_power_definition", _ma(x["P_ess"] - x["P_ess_dis"] + x["P_ess_cha"]), _mx(pmax))
    rec("storage_power_capacity", max(_pos(x["P_ess_cha"] - pmax), _pos(x["P_ess_dis"] - pmax)), _mx(pmax))
    rec("storage_binary_mode_gates", max(_pos(x["P_ess_cha"] - bigm * x["u_ess"]), _pos(x["P_ess_dis"] - bigm * (1 - x["u_ess"]))), bigm)
    rec("storage_energy_bounds", max(_pos(float(op["ess_soc_min"]) * cap - x["E_ess"]), _pos(x["E_ess"] - cap)), _mx(cap))
    eta = float(op["ess_efficiency"]); chronology = []; wrap = []; terminal = []
    for s in range(4):
        o = 24 * s
        chronology.append(x["E_ess"][:, o + 1:o + 24] - x["E_ess"][:, o:o + 23] - eta * x["P_ess_cha"][:, o:o + 23] + x["P_ess_dis"][:, o:o + 23] / eta)
        e0 = float(op["ess_soc_initial"]) * cap[:, 0]
        wrap.append(x["E_ess"][:, o] - e0 - eta * x["P_ess_cha"][:, o + 23] + x["P_ess_dis"][:, o + 23] / eta)
        terminal.append(x["E_ess"][:, o + 23] - e0)
    rec("storage_hour_1_to_23_chronology", max(_ma(a) for a in chronology), _mx(cap))
    rec("storage_hour24_wrap_to_hour1", max(_ma(a) for a in wrap), _mx(cap))
    rec("storage_terminal_E24", max(_ma(a) for a in terminal), _mx(cap))

    # SOP fixed capacities and balances.
    s1 = np.asarray(plan["sop_units_endpoint_1"], float) * float(plan["sop_unit_mva"]) / sbase
    s2 = np.asarray(plan["sop_units_endpoint_2"], float) * float(plan["sop_unit_mva"]) / sbase
    sop_bound = sop_loss_cone = sop_cap_cone = 0.0; lossfac = float(op["sop_loss_factor"])
    for k, sk in ((1, s1), (2, s2)):
        pp, qq, ll = x[f"P_sop_{k}"], x[f"Q_sop_{k}"], x[f"Pl_sop_{k}"]
        sop_bound = max(sop_bound, _pos(np.abs(pp) - sk[:, None]), _pos(-qq), _pos(qq - float(op["sop_q_upper_fraction"]) * sk[:, None]), _pos(-ll), _pos(ll - lossfac * sk[:, None]))
        sop_loss_cone = max(sop_loss_cone, _pos(np.hypot(pp, qq) - ll / lossfac))
        sop_cap_cone = max(sop_cap_cone, _pos(np.hypot(pp, qq) - sk[:, None]))
    rec("SOP_endpoint_bounds", sop_bound)
    rec("SOP_loss_cones", sop_loss_cone)
    rec("SOP_capacity_cones", sop_cap_cone)
    rec("SOP_active_loss_balance", _ma(x["P_sop_1"] + x["P_sop_2"] + x["Pl_sop_1"] + x["Pl_sop_2"]))

    # Shedding and nodal balances.
    rec("load_shedding_mask", max(_pos(-x["P_ls"]), _pos(x["P_ls"] - pl * shed_mask)), _mx(pl))
    rec("IDC_90_percent_service", _pos(x["P_ls"][49, fault_h] - 0.10 * pl[49, fault_h]), _mx(pl[49, fault_h]))
    station_b = np.asarray(sets["existing_substations"] + sets["candidate_substations"], int) - 1
    pv_b = np.asarray(sets["pv_buses"], int) - 1; wt_b = np.asarray(sets["wind_buses"], int) - 1
    sop_rows = np.asarray(sets["sop_line_rows"], int) - 1
    pbal = -afr @ x["P"] + ato @ x["P"] + ato @ (rr[:, None] * x["l"]) - (pl - x["P_ls"])
    qbal = -afr @ x["Q"] + ato @ x["Q"] + ato @ (xx[:, None] * x["l"]) - (ql - load_qfac * x["P_ls"])
    pbal[station_b] += x["P_s"]; qbal[station_b] += x["Q_s"]
    pbal[pv_b] += x["P_pv"]; qbal[pv_b] += x["Q_pv"]
    pbal[wt_b] += x["P_wt"]; qbal[wt_b] += x["Q_wt"]
    pbal[ess_buses] += x["P_ess"]
    pbal[fr[sop_rows]] += x["P_sop_1"]; qbal[fr[sop_rows]] += x["Q_sop_1"]
    pbal[to[sop_rows]] += x["P_sop_2"]; qbal[to[sop_rows]] += x["Q_sop_2"]
    rec("nodal_active_balance", _ma(pbal), max(_mx(pl), _mx(x["P_s"])))
    rec("nodal_reactive_balance", _ma(qbal), max(_mx(ql), _mx(x["Q_s"])))

    # Objective, independently reconstructed term by term.
    disc = float(plan["discount_rate"]); years = int(plan["years"])
    crf = disc * (1 + disc) ** years / ((1 + disc) ** years - 1)
    line_inv_coef = float(plan["line_build_cost_usd_per_km"]) * length.copy(); line_inv_coef[existing] = 0.0
    obj = {
        "investment_existing_station": float(np.sum(np.asarray(plan["existing_station_option_cost_usd"]) * x["uEX"])),
        "investment_candidate_station": float(np.sum(np.asarray(plan["candidate_station_option_cost_usd"]) * x["uNS"])),
        "investment_lines": float(np.dot(line_inv_coef, ul)),
        "investment_wind": float(econ["wind_investment_usd_per_mw"]) * float(plan["dg_unit_mw"]) * float(np.sum(x["zWTG"])),
        "investment_pv": float(econ["pv_investment_usd_per_mw"]) * float(plan["dg_unit_mw"]) * float(np.sum(x["zPVG"])),
        "investment_sop_fixed": float(econ["sop_investment_usd_per_mva"]) * float(plan["sop_unit_mva"]) * float(np.sum(plan["sop_units_endpoint_1"]) + np.sum(plan["sop_units_endpoint_2"])),
        "investment_storage": float(econ["ess_investment_usd_per_mwh"]) * sbase * float(np.dot(sess, x["zESS"][:, 0])),
    }
    obj["C_inv"] = sum(obj[k] for k in list(obj) if k.startswith("investment_"))
    obj["F_inv"] = crf * obj["C_inv"]
    obj["maintenance_pv"] = float(np.sum(x["zPVG"])) * float(plan["dg_unit_mw"]) * float(econ["pv_maintenance_usd_per_mwh"]) * float(econ["pv_capacity_factor"]) * 8760.0
    obj["maintenance_wind"] = float(np.sum(x["zWTG"])) * float(plan["dg_unit_mw"]) * float(econ["wind_maintenance_usd_per_mwh"]) * float(econ["wind_capacity_factor"]) * 8760.0
    obj["maintenance_sop_fixed"] = float(np.sum(plan["sop_units_endpoint_1"]) + np.sum(plan["sop_units_endpoint_2"])) * float(plan["sop_unit_mva"]) * float(econ["sop_maintenance_fraction"]) * float(econ["sop_investment_usd_per_mva"])
    obj["maintenance_lines"] = 0.25 * 4.0 * float(np.dot(float(plan["line_maintenance_usd_per_km_year"]) * length, ul))
    obj["C_ma"] = obj["maintenance_pv"] + obj["maintenance_wind"] + obj["maintenance_sop_fixed"] + obj["maintenance_lines"]
    price = np.tile(np.asarray(op["purchase_price_usd_per_mwh"], float), 4)
    obj["operating_purchase"] = sbase * float(np.sum(price[None, :] * x["P_s"]))
    wt_avail = x["zWTG"][:, 0, None] * fc[:, 0][None, :] * sdg
    pv_avail = x["zPVG"][:, 0, None] * fc[:, 1][None, :] * sdg
    obj["operating_wind_curtailment"] = sbase * float(econ["curtailment_cost_usd_per_mwh"]) * float(np.sum(wt_avail - x["P_wt"]))
    obj["operating_pv_curtailment"] = sbase * float(econ["curtailment_cost_usd_per_mwh"]) * float(np.sum(pv_avail - x["P_pv"]))
    obj["operating_load_shed"] = sbase * float(econ["load_shed_cost_usd_per_mwh"]) * float(np.sum(x["P_ls"]))
    obj["C_ope"] = obj["operating_purchase"] + obj["operating_wind_curtailment"] + obj["operating_pv_curtailment"] + obj["operating_load_shed"]
    obj["F_ope"] = float(econ["scenario_day_annualization_factor"]) * obj["C_ope"] + obj["C_ma"]
    obj["full_reported_model_total"] = obj["F_inv"] + obj["F_ope"]
    obj["fixed_SOP_total_contribution"] = crf * obj["investment_sop_fixed"] + obj["maintenance_sop_fixed"]
    obj["solver_objective_if_fixed_SOP_constants_dropped"] = obj["full_reported_model_total"] - obj["fixed_SOP_total_contribution"]

    violated = [name for name, val in families.items() if val["normalized"] > tolerance]
    passed = not structural_errors and not violated and max_integrality <= tolerance
    return {
        "passed": passed,
        "replay_audit_passed": passed,
        "same_model_replay_passed": passed,
        "construction_model_imported": False,
        "authoritative_inputs": ["line_load.xlsx", "model_spec.json"],
        "tolerance": tolerance,
        "solution_path": str(solution_path),
        "structural_errors": structural_errors,
        "maximum_integrality_distance": max_integrality,
        "integrality_by_family": integ_by,
        "maximum_normalized_residual": max((v["normalized"] for v in families.values()), default=0.0),
        "violated_families": violated,
        "families": families,
        "objective_components": obj,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", type=Path, default=Path("/work/app"))
    ap.add_argument("--solution", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--tolerance", type=float, default=5e-5)
    args = ap.parse_args()
    result = audit(args.root, args.solution, args.tolerance)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")
    print(json.dumps({k: result.get(k) for k in ("passed", "maximum_normalized_residual", "maximum_integrality_distance", "violated_families")}, indent=2))
    raise SystemExit(0 if result.get("passed") else 2)


if __name__ == "__main__":
    main()
