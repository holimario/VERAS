"""Input loading and deterministic derived data for model construction only."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl


def load_inputs(root: Path) -> dict:
    spec = json.loads((root / "model_spec.json").read_text())
    wb = openpyxl.load_workbook(root / "line_load.xlsx", data_only=True, read_only=True)
    ws1 = wb["Sheet1"]
    ws3 = wb["Sheet3"]

    branch = np.asarray(
        [[ws1.cell(r, c).value for c in range(1, 5)] for r in range(2, 65)],
        dtype=float,
    )
    bus = np.asarray(
        [[ws1.cell(r, c).value for c in range(7, 9)] for r in range(2, 56)],
        dtype=float,
    )
    forecast = np.asarray(
        [[ws3.cell(r, c).value for c in range(2, 5)] for r in range(2, 98)],
        dtype=float,
    )
    wb.close()

    if branch.shape != (63, 4) or bus.shape != (54, 2) or forecast.shape != (96, 3):
        raise ValueError("authoritative workbook ranges have unexpected shapes")
    if not np.array_equal(branch[:, 0], np.arange(1, 64)):
        raise ValueError("branch identifiers are not 1..63 in workbook order")
    if not np.array_equal(bus[:, 0], np.arange(1, 55)):
        raise ValueError("bus identifiers are not 1..54 in workbook order")

    net = spec["network"]
    plan = spec["planning"]
    op = spec["operation"]
    sets = spec["sets"]
    n_b, n_e, n_h, n_s = 54, 63, 96, 4

    fr = branch[:, 1].astype(int) - 1
    to = branch[:, 2].astype(int) - 1
    length = branch[:, 3]
    a_fr = np.zeros((n_b, n_e))
    a_to = np.zeros((n_b, n_e))
    a_fr[fr, np.arange(n_e)] = 1.0
    a_to[to, np.arange(n_e)] = 1.0

    s_base = float(net["s_base_mva"])
    v_base = float(net["v_base_kv"])
    i_base = s_base / v_base
    z_base = v_base * v_base / s_base
    i_max = float(net["line_current_limit_ka"]) / i_base
    i_sq_max = i_max * i_max
    s_line_max = i_max * float(net["v_max_pu"])
    resistance = float(net["resistance_ohm_per_km"]) * length / z_base
    reactance = float(net["reactance_ohm_per_km"]) * length / z_base

    wt_fc = forecast[:, 0]
    pv_fc = forecast[:, 1]
    pl_fc = forecast[:, 2]
    pl = bus[:, 1, None] * pl_fc[None, :] / s_base
    load_qfac = np.tan(np.arccos(float(net["load_power_factor"])))
    ql = pl * load_qfac
    dg_qfac = np.sqrt(1.0 - float(net["dg_min_power_factor"]) ** 2) / float(
        net["dg_min_power_factor"]
    )

    mask = np.ones((n_e, n_h))
    fault_e = np.asarray(sets["fault_line_rows"], dtype=int) - 1
    fault_h = np.asarray(spec["fault"]["fault_hours_one_based"], dtype=int) - 1
    mask[np.ix_(fault_e, fault_h)] = 0.0
    shed_mask = np.zeros((n_b, n_h))
    shed_mask[:, fault_h] = 1.0

    price24 = np.asarray(op["purchase_price_usd_per_mwh"], dtype=float)
    price96 = np.tile(price24, n_s)
    ess_buses = np.asarray(sets["ess_buses"], dtype=int) - 1
    sess_rated = np.full(len(ess_buses), float(op["ess_unit_mwh_non_idc"]) / s_base)
    sess_rated[ess_buses == (sets["idc_buses"][0] - 1)] = (
        float(op["ess_unit_mwh_idc"]) / s_base
    )
    sdg_rated = float(plan["dg_unit_mw"]) / s_base

    r = float(plan["discount_rate"])
    years = int(plan["years"])
    crf = r * (1.0 + r) ** years / ((1.0 + r) ** years - 1.0)

    return {
        "spec": spec,
        "branch": branch,
        "bus": bus,
        "forecast": forecast,
        "fr": fr,
        "to": to,
        "length": length,
        "A_fr": a_fr,
        "A_to": a_to,
        "A": a_to - a_fr,
        "R": resistance,
        "X": reactance,
        "I_sq_max": i_sq_max,
        "S_line_max": s_line_max,
        "wt_fc": wt_fc,
        "pv_fc": pv_fc,
        "pl_fc": pl_fc,
        "PL": pl,
        "QL": ql,
        "load_qfac": load_qfac,
        "dg_qfac": dg_qfac,
        "mask": mask,
        "shed_mask": shed_mask,
        "price96": price96,
        "ess_buses": ess_buses,
        "pv_buses": np.asarray(sets["pv_buses"], dtype=int) - 1,
        "wt_buses": np.asarray(sets["wind_buses"], dtype=int) - 1,
        "station_buses": np.asarray(sets["existing_substations"] + sets["candidate_substations"], dtype=int) - 1,
        "sop_rows": np.asarray(sets["sop_line_rows"], dtype=int) - 1,
        "existing_rows": np.asarray(sets["existing_line_rows"], dtype=int) - 1,
        "fault_hours": fault_h,
        "Sess_rated": sess_rated,
        "Sdg_rated": sdg_rated,
        "crf": crf,
    }
