#!/usr/bin/env python3
"""Independent public-data replay for the submitted ADN solution artifact."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Dict, Iterable, Mapping, Tuple

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"

REQUIRED_SHAPES = {
    "uL": (63, 1), "uNS": (2, 2), "uEX": (2, 2),
    "zESS": (7, 1), "zPVG": (7, 1), "zWTG": (6, 1),
    "u_ess": (7, 96), "kt": (4, 96), "FF": (63, 4),
    "FS": (54, 4), "P": (63, 96), "Q": (63, 96),
    "l": (63, 96), "v": (54, 96), "P_s": (4, 96),
    "Q_s": (4, 96), "P_pv": (7, 96), "Q_pv": (7, 96),
    "P_wt": (6, 96), "Q_wt": (6, 96), "E_ess": (7, 96),
    "P_ess": (7, 96), "P_ess_cha": (7, 96),
    "P_ess_dis": (7, 96), "P_sop_1": (4, 96),
    "P_sop_2": (4, 96), "Q_sop_1": (4, 96),
    "Q_sop_2": (4, 96), "Pl_sop_1": (4, 96),
    "Pl_sop_2": (4, 96), "P_ls": (54, 96),
}


def load_public_data() -> Dict[str, np.ndarray | dict | float]:
    with (DATA_DIR / "model_spec.json").open("r", encoding="utf-8") as f:
        spec = json.load(f)

    wb = openpyxl.load_workbook(
        DATA_DIR / "line_load.xlsx", data_only=True, read_only=True
    )
    ws = wb["Sheet1"]
    wf = wb["Sheet3"]
    branch = np.asarray(
        [[ws.cell(r, c).value for c in range(1, 5)] for r in range(2, 65)],
        dtype=float,
    )
    bus = np.asarray(
        [[ws.cell(r, c).value for c in range(7, 9)] for r in range(2, 56)],
        dtype=float,
    )
    forecast = np.asarray(
        [[wf.cell(r, c).value for c in range(2, 5)] for r in range(2, 98)],
        dtype=float,
    )
    wb.close()

    if not np.array_equal(branch[:, 0], np.arange(1, 64)):
        raise ValueError("Workbook branch IDs are not exactly 1..63")
    if not np.array_equal(bus[:, 0], np.arange(1, 55)):
        raise ValueError("Workbook bus IDs are not exactly 1..54")

    frm = branch[:, 1].astype(int) - 1
    to = branch[:, 2].astype(int) - 1
    length = branch[:, 3]
    base_load = bus[:, 1]
    wt_fc, pv_fc, pl_fc = forecast.T

    net = spec["network"]
    sbase = float(net["s_base_mva"])
    vbase = float(net["v_base_kv"])
    ibase = sbase / vbase
    zbase = vbase * vbase / sbase
    imax = float(net["line_current_limit_ka"]) / ibase
    vmax = float(net["v_max_pu"])
    r = float(net["resistance_ohm_per_km"]) * length / zbase
    x = float(net["reactance_ohm_per_km"]) * length / zbase
    pf_load = float(net["load_power_factor"])
    tau = math.tan(math.acos(pf_load))
    pl = base_load[:, None] * pl_fc[None, :] / sbase
    ql = tau * pl

    return {
        "spec": spec, "frm": frm, "to": to, "length": length,
        "base_load": base_load, "wt_fc": wt_fc, "pv_fc": pv_fc,
        "pl_fc": pl_fc, "sbase": sbase, "r": r, "x": x,
        "imax": imax, "imax2": imax * imax,
        "sline": imax * vmax, "tau": tau, "PL": pl, "QL": ql,
    }


class ResidualBook:
    def __init__(self) -> None:
        self.groups: Dict[str, Dict[str, float | str]] = {}

    @staticmethod
    def _finite_max(values: Iterable[float] | np.ndarray) -> float:
        a = np.asarray(values, dtype=float)
        return 0.0 if a.size == 0 else float(np.max(a))

    def eq(self, name: str, residual: np.ndarray | float, scale: float = 1.0) -> None:
        raw = self._finite_max(np.abs(residual))
        self.groups[name] = {
            "type": "equality", "max_residual": raw,
            "max_normalized_residual": raw / max(1.0, float(scale)),
        }

    def le(self, name: str, violation: np.ndarray | float, scale: float = 1.0) -> None:
        raw = self._finite_max(np.maximum(np.asarray(violation, dtype=float), 0.0))
        self.groups[name] = {
            "type": "inequality", "max_residual": raw,
            "max_normalized_residual": raw / max(1.0, float(scale)),
        }

    def merge_max(self, name: str, kind: str, values: Iterable[np.ndarray | float], scale: float = 1.0) -> None:
        vals = [np.asarray(v, dtype=float).ravel() for v in values]
        a = np.concatenate(vals) if vals else np.zeros(0)
        if kind == "eq":
            self.eq(name, a, scale)
        else:
            self.le(name, a, scale)


def _objective(data: Mapping, a: Mapping[str, np.ndarray]) -> Tuple[float, Dict[str, float]]:
    spec = data["spec"]
    plan = spec["planning"]
    econ = spec["economics"]
    op = spec["operation"]
    length = data["length"]
    sbase = float(data["sbase"])

    uL = a["uL"][:, 0]
    uEX, uNS = a["uEX"], a["uNS"]
    zess = a["zESS"][:, 0]
    zpv = a["zPVG"][:, 0]
    zwt = a["zWTG"][:, 0]
    existing = np.asarray(spec["sets"]["existing_line_rows"], dtype=int) - 1
    line_cost = float(plan["line_build_cost_usd_per_km"]) * length
    line_cost[existing] = 0.0
    ess_rated = np.full(7, float(op["ess_unit_mwh_non_idc"]) / sbase)
    ess_rated[-1] = float(op["ess_unit_mwh_idc"]) / sbase
    zsop1 = np.asarray(plan["sop_units_endpoint_1"], dtype=float)
    zsop2 = np.asarray(plan["sop_units_endpoint_2"], dtype=float)

    station_inv = float(np.sum(np.asarray(plan["existing_station_option_cost_usd"]) * uEX))
    station_inv += float(np.sum(np.asarray(plan["candidate_station_option_cost_usd"]) * uNS))
    line_inv = float(line_cost @ uL)
    wind_inv = float(econ["wind_investment_usd_per_mw"]) * 0.1 * float(np.sum(zwt))
    pv_inv = float(econ["pv_investment_usd_per_mw"]) * 0.1 * float(np.sum(zpv))
    sop_inv = float(econ["sop_investment_usd_per_mva"]) * 0.1 * float(np.sum(zsop1 + zsop2))
    ess_inv = float(econ["ess_investment_usd_per_mwh"]) * sbase * float(ess_rated @ zess)
    c_inv = station_inv + line_inv + wind_inv + pv_inv + sop_inv + ess_inv
    rate = float(plan["discount_rate"])
    years = int(plan["years"])
    crf = rate * (1.0 + rate) ** years / ((1.0 + rate) ** years - 1.0)
    f_inv = crf * c_inv

    pv_ma = 0.1 * float(econ["pv_maintenance_usd_per_mwh"]) * float(econ["pv_capacity_factor"]) * 8760.0 * float(np.sum(zpv))
    wt_ma = 0.1 * float(econ["wind_maintenance_usd_per_mwh"]) * float(econ["wind_capacity_factor"]) * 8760.0 * float(np.sum(zwt))
    sop_ma = 0.1 * float(econ["sop_maintenance_fraction"]) * float(econ["sop_investment_usd_per_mva"]) * float(np.sum(zsop1 + zsop2))
    line_ma = 0.25 * 4.0 * float(plan["line_maintenance_usd_per_km_year"]) * float(length @ uL)
    c_ma = pv_ma + wt_ma + sop_ma + line_ma

    price = np.tile(np.asarray(op["purchase_price_usd_per_mwh"], dtype=float), 4)
    purchase = sbase * float(np.sum(a["P_s"] * price[None, :]))
    sdg = float(plan["dg_unit_mw"]) / sbase
    wt_av = zwt[:, None] * data["wt_fc"][None, :] * sdg
    pv_av = zpv[:, None] * data["pv_fc"][None, :] * sdg
    curtail = sbase * float(econ["curtailment_cost_usd_per_mwh"]) * (
        float(np.sum(wt_av - a["P_wt"])) + float(np.sum(pv_av - a["P_pv"]))
    )
    shed = sbase * float(econ["load_shed_cost_usd_per_mwh"]) * float(np.sum(a["P_ls"]))
    c_ope = purchase + curtail + shed
    f_ope = float(econ["scenario_day_annualization_factor"]) * c_ope + c_ma
    total = f_inv + f_ope
    return total, {
        "station_investment": station_inv, "line_investment": line_inv,
        "wind_investment": wind_inv, "pv_investment": pv_inv,
        "sop_investment_constant": sop_inv, "storage_investment": ess_inv,
        "capital_recovery_factor": crf, "annualized_investment": f_inv,
        "pv_maintenance": pv_ma, "wind_maintenance": wt_ma,
        "sop_maintenance_constant": sop_ma, "line_maintenance": line_ma,
        "purchase_96_columns": purchase, "curtailment_96_columns": curtail,
        "load_shed_96_columns": shed, "operating_96_columns": c_ope,
        "annual_operating_and_maintenance": f_ope, "total_objective": total,
    }


def audit_arrays(data: Mapping, arrays: Mapping[str, np.ndarray], tol: float = 5e-5) -> Dict:
    missing = sorted(set(REQUIRED_SHAPES) - set(arrays))
    extra = sorted(set(arrays) - set(REQUIRED_SHAPES))
    shape_errors = {
        k: {"expected": REQUIRED_SHAPES[k], "actual": tuple(np.shape(arrays[k]))}
        for k in REQUIRED_SHAPES if k in arrays and tuple(np.shape(arrays[k])) != REQUIRED_SHAPES[k]
    }
    nonfinite = {
        k: int(np.size(arrays[k]) - np.count_nonzero(np.isfinite(arrays[k])))
        for k in REQUIRED_SHAPES if k in arrays
        and tuple(np.shape(arrays[k])) == REQUIRED_SHAPES[k]
        and not np.all(np.isfinite(arrays[k]))
    }
    dtype_errors = {
        k: str(np.asarray(arrays[k]).dtype)
        for k in REQUIRED_SHAPES if k in arrays
        and np.asarray(arrays[k]).dtype != np.dtype("float64")
    }
    if missing or extra or shape_errors or nonfinite or dtype_errors:
        return {
            "feasible": False, "schema_ok": False, "missing": missing,
            "extra": extra, "shape_errors": shape_errors,
            "nonfinite": nonfinite, "dtype_errors": dtype_errors, "groups": {},
            "max_equality_residual": math.inf,
            "max_inequality_residual": math.inf,
            "max_normalized_constraint_residual": math.inf,
            "max_integrality_residual": math.inf,
            "violated_constraint_groups": ["schema"],
            "number_violated_constraint_groups": 1,
            "decision_universe_coverage": f"{len(REQUIRED_SHAPES) - len(missing)}/{len(REQUIRED_SHAPES)}",
            "objective": None,
        }

    a = {k: np.asarray(arrays[k], dtype=float) for k in REQUIRED_SHAPES}
    spec = data["spec"]
    net, plan, op = spec["network"], spec["planning"], spec["operation"]
    bset = spec["sets"]
    B, E, H, S = 54, 63, 96, 4
    M = float(net["big_m"])
    vmin2, vmax2 = float(net["v_min_pu"]) ** 2, float(net["v_max_pu"]) ** 2
    frm, to = data["frm"], data["to"]
    r, x, PL, QL = data["r"], data["x"], data["PL"], data["QL"]
    tau, imax2, sline = float(data["tau"]), float(data["imax2"]), float(data["sline"])
    sdg = float(plan["dg_unit_mw"]) / float(data["sbase"])
    qfac = math.sqrt(1.0 - float(net["dg_min_power_factor"]) ** 2) / float(net["dg_min_power_factor"])

    uL, uNS, uEX = a["uL"][:, 0], a["uNS"], a["uEX"]
    zess, zpv, zwt = a["zESS"][:, 0], a["zPVG"][:, 0], a["zWTG"][:, 0]
    rb = ResidualBook()

    binary_names = ["uL", "uNS", "uEX", "u_ess"]
    integer_names = binary_names + ["zESS", "zPVG", "zWTG", "kt"]
    int_res = max(float(np.max(np.abs(a[k] - np.rint(a[k])))) for k in integer_names)
    integrality_exact = all(np.array_equal(a[k], np.rint(a[k])) for k in integer_names)
    rb.merge_max("binary_domains", "le", [
        -a[k] for k in binary_names
    ] + [a[k] - 1.0 for k in binary_names])
    rb.merge_max("integer_count_bounds", "le", [
        1.0 - zess, zess - 20.0, 1.0 - zpv, zpv - 50.0,
        1.0 - zwt, zwt - 50.0, -8.0 - a["kt"], a["kt"] - 8.0,
    ])

    existing = np.asarray(bset["existing_line_rows"], dtype=int) - 1
    rb.eq("existing_line_fixing", uL[existing] - 1.0)
    rb.le("planning_count_limits", np.asarray([
        np.sum(uL) - (52.0 - np.sum(uNS)),
        np.sum(zess) - float(plan["ess_units_total_max"]),
        np.sum(zpv) - float(plan["pv_units_total_max"]),
        np.max(np.sum(uEX, axis=1) - 1.0),
        np.max(np.sum(uNS, axis=1) - 1.0),
    ]), scale=63.0)
    rb.eq("topology_cardinality", np.sum(uL) - (52.0 - np.sum(uNS)), scale=54.0)

    ff, fs = a["FF"], a["FS"]
    rb.merge_max("artificial_flow_bounds", "le", [ff - 54.0 * uL[:, None], -ff - 54.0 * uL[:, None]], 54.0)
    topo_bal = np.zeros((B, S))
    for e in range(E):
        topo_bal[to[e], :] += ff[e, :]
        topo_bal[frm[e], :] -= ff[e, :]
    topo_bal += fs - 1.0
    rb.eq("artificial_flow_balance", topo_bal, 54.0)
    rb.le("artificial_root_nonnegativity", -fs, 54.0)
    root_v = [np.abs(fs[:50, :])]
    root_v += [fs[50, :] - 54.0 * np.sum(uEX[0, :])]
    root_v += [fs[51, :] - 54.0 * np.sum(uEX[1, :])]
    root_v += [fs[52, :] - 54.0 * np.sum(uNS[0, :])]
    root_v += [fs[53, :] - 54.0 * np.sum(uNS[1, :])]
    rb.merge_max("artificial_root_gates", "le", root_v, 54.0)

    cap = np.asarray([
        1.67 + 1.33 * uEX[0, 0] + 1.67 * uEX[0, 1],
        1.67 + 1.33 * uEX[1, 0] + 1.67 * uEX[1, 1],
        1.67 * uNS[0, 0] + 2.22 * uNS[0, 1],
        1.67 * uNS[1, 0] + 2.22 * uNS[1, 1],
    ])
    derate = np.ones((4, H)); derate[:2, 6:8] = 0.30
    station_norm = np.hypot(a["P_s"], a["Q_s"])
    rb.merge_max("station_domains_and_cones", "le", [
        -a["P_s"], station_norm - 0.5 * cap[:, None] * derate
    ], 2.5)

    rb.merge_max("ordinary_bus_voltage_bounds", "le", [
        vmin2 - a["v"][:50, :], a["v"][:50, :] - vmax2
    ], vmax2)
    rb.eq("existing_station_voltage_taps", np.vstack([
        a["v"][50, :] - 1.01 - 0.025 * a["kt"][0, :],
        a["v"][51, :] - 1.01 - 0.025 * a["kt"][1, :],
    ]), vmax2)
    new_on = np.sum(uNS, axis=1)
    cand_v = []
    for i in range(2):
        vv, kk, on = a["v"][52 + i, :], a["kt"][2 + i, :], new_on[i]
        tapv = 1.01 + 0.025 * kk
        cand_v.extend([
            tapv - M * (1.0 - on) - vv,
            vv - tapv - M * (1.0 - on),
            vmin2 - M * on - vv,
            vv - vmax2 - M * on,
            -8.0 * on - kk,
            kk - 8.0 * on,
        ])
    rb.merge_max("candidate_station_voltage_gates", "le", cand_v, M)

    mask = np.ones((E, H)); mask[[33, 43], 6:8] = 0.0
    yeff = uL[:, None] * mask
    rb.merge_max("branch_bounds_and_fault_mask", "le", [
        -a["l"], a["l"] - imax2 * yeff,
        a["P"] - sline * yeff, -a["P"] - sline * yeff,
        a["Q"] - sline * yeff, -a["Q"] - sline * yeff,
    ], sline)
    lhs_soc = np.sqrt(
        (2.0 * a["P"]) ** 2 + (2.0 * a["Q"]) ** 2
        + (a["v"][frm, :] - a["l"]) ** 2
    )
    rhs_soc = a["v"][frm, :] + a["l"]
    rb.le("branch_soc", lhs_soc - rhs_soc, 2.0)
    drop = a["v"][frm, :] - a["v"][to, :] - 2.0 * r[:, None] * a["P"] - 2.0 * x[:, None] * a["Q"] + (r * r + x * x)[:, None] * a["l"]
    rb.merge_max("branch_voltage_drop_gates", "le", [
        drop - M * (1.0 - yeff), -drop - M * (1.0 - yeff)
    ], M)

    for tag, z, fc, pg, qg in [
        ("wind", zwt, data["wt_fc"], a["P_wt"], a["Q_wt"]),
        ("pv", zpv, data["pv_fc"], a["P_pv"], a["Q_pv"]),
    ]:
        avail = z[:, None] * fc[None, :] * sdg
        rb.merge_max(f"{tag}_dispatch_bounds", "le", [0.3 * avail - pg, pg - avail])
        rb.le(f"{tag}_installed_capacity_cone", np.hypot(pg, qg) - z[:, None] * sdg)
        rb.merge_max(f"{tag}_actual_power_factor", "le", [qg - qfac * pg, -qg - qfac * pg])

    sess = np.full(7, float(op["ess_unit_mwh_non_idc"]) / float(data["sbase"]))
    sess[-1] = float(op["ess_unit_mwh_idc"]) / float(data["sbase"])
    ecap = zess * sess; pmax = float(op["ess_power_to_energy_ratio"]) * ecap
    rb.eq("storage_net_power", a["P_ess"] - a["P_ess_dis"] + a["P_ess_cha"])
    rb.merge_max("storage_power_and_mode", "le", [
        -a["P_ess_cha"], a["P_ess_cha"] - pmax[:, None],
        -a["P_ess_dis"], a["P_ess_dis"] - pmax[:, None],
        a["P_ess_cha"] - M * a["u_ess"],
        a["P_ess_dis"] - M * (1.0 - a["u_ess"]),
    ])
    rb.merge_max("storage_energy_bounds", "le", [
        float(op["ess_soc_min"]) * ecap[:, None] - a["E_ess"],
        a["E_ess"] - ecap[:, None],
    ])
    chron = []
    eta = float(op["ess_efficiency"])
    e0 = float(op["ess_soc_initial"]) * ecap
    for s in range(4):
        o = 24 * s
        chron.append(a["E_ess"][:, o + 1:o + 24] - a["E_ess"][:, o:o + 23] - eta * a["P_ess_cha"][:, o:o + 23] + a["P_ess_dis"][:, o:o + 23] / eta)
        chron.append(a["E_ess"][:, o] - e0 - eta * a["P_ess_cha"][:, o + 23] + a["P_ess_dis"][:, o + 23] / eta)
        chron.append(a["E_ess"][:, o + 23] - e0)
    rb.merge_max("storage_wraparound_chronology", "eq", chron)

    zs1 = np.asarray(plan["sop_units_endpoint_1"], dtype=float)
    zs2 = np.asarray(plan["sop_units_endpoint_2"], dtype=float)
    sopunit = float(plan["sop_unit_mva"]) / float(data["sbase"])
    for tag, ss, pp, qq, loss in [
        ("endpoint_1", zs1 * sopunit, a["P_sop_1"], a["Q_sop_1"], a["Pl_sop_1"]),
        ("endpoint_2", zs2 * sopunit, a["P_sop_2"], a["Q_sop_2"], a["Pl_sop_2"]),
    ]:
        rb.merge_max(f"sop_{tag}_bounds", "le", [
            pp - ss[:, None], -pp - ss[:, None], -qq,
            qq - float(op["sop_q_upper_fraction"]) * ss[:, None],
            -loss, loss - float(op["sop_loss_factor"]) * ss[:, None],
        ])
        norm = np.hypot(pp, qq)
        rb.le(f"sop_{tag}_loss_cone", norm - loss / float(op["sop_loss_factor"]))
        rb.le(f"sop_{tag}_capacity_cone", norm - ss[:, None])
    rb.eq("sop_active_loss_balance", a["P_sop_1"] + a["P_sop_2"] + a["Pl_sop_1"] + a["Pl_sop_2"])

    shedmask = np.zeros((B, H)); shedmask[:, 6:8] = 1.0
    rb.merge_max("load_shedding_mask", "le", [-a["P_ls"], a["P_ls"] - PL * shedmask])
    rb.le("idc_fault_service", a["P_ls"][49, 6:8] - (1.0 - float(op["idc_min_served_fraction_fault"])) * PL[49, 6:8])

    pbal = -(PL - a["P_ls"])
    qbal = -(QL - tau * a["P_ls"])
    for e in range(E):
        pbal[frm[e], :] -= a["P"][e, :]
        pbal[to[e], :] += a["P"][e, :] + r[e] * a["l"][e, :]
        qbal[frm[e], :] -= a["Q"][e, :]
        qbal[to[e], :] += a["Q"][e, :] + x[e] * a["l"][e, :]
    for rr, b in enumerate([50, 51, 52, 53]):
        pbal[b, :] += a["P_s"][rr, :]; qbal[b, :] += a["Q_s"][rr, :]
    for ii, b in enumerate(np.asarray(bset["wind_buses"], dtype=int) - 1):
        pbal[b, :] += a["P_wt"][ii, :]; qbal[b, :] += a["Q_wt"][ii, :]
    for ii, b in enumerate(np.asarray(bset["pv_buses"], dtype=int) - 1):
        pbal[b, :] += a["P_pv"][ii, :]; qbal[b, :] += a["Q_pv"][ii, :]
    for ii, b in enumerate(np.asarray(bset["ess_buses"], dtype=int) - 1):
        pbal[b, :] += a["P_ess"][ii, :]
    sop_rows = np.asarray(bset["sop_line_rows"], dtype=int) - 1
    for mm, e in enumerate(sop_rows):
        pbal[frm[e], :] += a["P_sop_1"][mm, :]
        qbal[frm[e], :] += a["Q_sop_1"][mm, :]
        pbal[to[e], :] += a["P_sop_2"][mm, :]
        qbal[to[e], :] += a["Q_sop_2"][mm, :]
    rb.eq("nodal_active_balance", pbal, 5.0)
    rb.eq("nodal_reactive_balance", qbal, 5.0)

    objective, objective_components = _objective(data, a)
    eqvals = [float(g["max_residual"]) for g in rb.groups.values() if g["type"] == "equality"]
    ineqvals = [float(g["max_residual"]) for g in rb.groups.values() if g["type"] == "inequality"]
    normvals = [float(g["max_normalized_residual"]) for g in rb.groups.values()]
    maxeq = max(eqvals, default=0.0); maxineq = max(ineqvals, default=0.0)
    theta = max(normvals, default=0.0)
    violated = sorted([name for name, g in rb.groups.items() if float(g["max_residual"]) > tol])
    if not integrality_exact:
        violated.append("integrality_exact")
        violated.sort()
    feasible = maxeq <= tol and maxineq <= tol and integrality_exact
    return {
        "feasible": bool(feasible), "schema_ok": True, "missing": missing,
        "extra": extra, "shape_errors": shape_errors, "nonfinite": nonfinite,
        "dtype_errors": dtype_errors,
        "groups": rb.groups, "max_equality_residual": maxeq,
        "max_inequality_residual": maxineq,
        "max_normalized_constraint_residual": theta,
        "max_integrality_residual": int_res,
        "integrality_exact": integrality_exact,
        "violated_constraint_groups": violated,
        "number_violated_constraint_groups": len(violated),
        "decision_universe_coverage": f"{len(REQUIRED_SHAPES)}/{len(REQUIRED_SHAPES)}",
        "objective": objective, "objective_components": objective_components,
        "tolerance": tol,
    }


def audit_npz(path: str | Path, tol: float = 5e-5) -> Dict:
    data = load_public_data()
    with np.load(path, allow_pickle=False) as z:
        arrays = {k: z[k] for k in z.files}
    return audit_arrays(data, arrays, tol=tol)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("npz", nargs="?", default=str(ROOT / "submissions" / "solution.npz"))
    parser.add_argument("--tol", type=float, default=5e-5)
    args = parser.parse_args()
    print(json.dumps(audit_npz(args.npz, args.tol), indent=2, sort_keys=True))
