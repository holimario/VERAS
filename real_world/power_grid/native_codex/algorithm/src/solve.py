#!/usr/bin/env python3
"""Build and solve the exact neutral-specification MISOCP with MOSEK Fusion."""

from __future__ import annotations

import argparse
import json
import math
import os
import time
from pathlib import Path

import numpy as np
import openpyxl
from mosek.fusion import AccSolutionStatus, Domain, Expr, Model, ObjectiveSense, SolutionStatus


def add_terms(terms):
    """Return a scalar Fusion expression for a sequence of scalar terms."""
    terms = list(terms)
    if not terms:
        return 0.0
    out = terms[0]
    for term in terms[1:]:
        out = Expr.add(out, term)
    return out


def le(model, expression, rhs=0.0):
    model.constraint(expression, Domain.lessThan(float(rhs)))


def ge(model, expression, rhs=0.0):
    model.constraint(expression, Domain.greaterThan(float(rhs)))


def eq(model, expression, rhs=0.0):
    model.constraint(expression, Domain.equalsTo(float(rhs)))


def read_inputs(input_root: Path):
    with (input_root / "model_spec.json").open("r", encoding="utf-8") as handle:
        spec = json.load(handle)

    workbook = openpyxl.load_workbook(
        input_root / "line_load.xlsx", read_only=True, data_only=True
    )
    ws1 = workbook["Sheet1"]
    ws3 = workbook["Sheet3"]
    branch = np.asarray(
        [[ws1.cell(row, col).value for col in range(1, 5)] for row in range(2, 65)],
        dtype=np.float64,
    )
    bus = np.asarray(
        [[ws1.cell(row, col).value for col in range(7, 9)] for row in range(2, 56)],
        dtype=np.float64,
    )
    forecast = np.asarray(
        [[ws3.cell(row, col).value for col in range(2, 5)] for row in range(2, 98)],
        dtype=np.float64,
    )
    workbook.close()

    if branch.shape != (63, 4) or bus.shape != (54, 2) or forecast.shape != (96, 3):
        raise ValueError("Workbook ranges do not have the required dimensions")
    if not np.array_equal(branch[:, 0], np.arange(1, 64, dtype=np.float64)):
        raise ValueError("Branch identifiers are not 1..63 in workbook row order")
    if not np.array_equal(bus[:, 0], np.arange(1, 55, dtype=np.float64)):
        raise ValueError("Bus identifiers are not 1..54 in workbook row order")
    return spec, branch, bus, forecast


def build_model(
    spec,
    branch,
    bus,
    forecast,
    log_path: Path,
    time_limit: float,
    threads: int,
    relax_integrality: bool = False,
    integer_fix: dict | None = None,
    exact_integer_names: set[str] | None = None,
):
    n_b, n_e, n_h, n_s = 54, 63, 96, 4
    network = spec["network"]
    planning = spec["planning"]
    operation = spec["operation"]
    economics = spec["economics"]

    big_m = float(network["big_m"])
    s_base = float(network["s_base_mva"])
    v_base = float(network["v_base_kv"])
    v_min_sq = float(network["v_min_pu"]) ** 2
    v_max_sq = float(network["v_max_pu"]) ** 2
    i_base = s_base / v_base
    i_max = float(network["line_current_limit_ka"]) / i_base
    i_sq_max = i_max**2
    s_line_max = i_max * float(network["v_max_pu"])
    z_base = v_base**2 / s_base
    lengths = branch[:, 3]
    fr = branch[:, 1].astype(int) - 1
    to = branch[:, 2].astype(int) - 1
    resistance = float(network["resistance_ohm_per_km"]) * lengths / z_base
    reactance = float(network["reactance_ohm_per_km"]) * lengths / z_base

    wt_fc = forecast[:, 0]
    pv_fc = forecast[:, 1]
    pl_fc = forecast[:, 2]
    pload = bus[:, 1, None] * pl_fc[None, :] / s_base
    tan_phi = math.tan(math.acos(float(network["load_power_factor"])))
    qload = pload * tan_phi
    qfac = math.sqrt(1.0 - float(network["dg_min_power_factor"]) ** 2) / float(
        network["dg_min_power_factor"]
    )
    sdg_rated = float(planning["dg_unit_mw"]) / s_base
    ess_unit_mwh = np.full(7, float(operation["ess_unit_mwh_non_idc"]), dtype=float)
    ess_unit_mwh[-1] = float(operation["ess_unit_mwh_idc"])
    ess_unit_pu = ess_unit_mwh / s_base
    eta = float(operation["ess_efficiency"])

    existing_lines = {int(e) - 1 for e in spec["sets"]["existing_line_rows"]}
    fault_lines = {int(e) - 1 for e in spec["sets"]["fault_line_rows"]}
    wt_buses = [int(b) - 1 for b in spec["sets"]["wind_buses"]]
    pv_buses = [int(b) - 1 for b in spec["sets"]["pv_buses"]]
    ess_buses = [int(b) - 1 for b in spec["sets"]["ess_buses"]]
    sop_lines = [int(e) - 1 for e in spec["sets"]["sop_line_rows"]]
    station_buses = [50, 51, 52, 53]

    outgoing = [[] for _ in range(n_b)]
    incoming = [[] for _ in range(n_b)]
    for e in range(n_e):
        outgoing[fr[e]].append(e)
        incoming[to[e]].append(e)

    mask = np.ones((n_e, n_h), dtype=float)
    for e in fault_lines:
        mask[e, 6:8] = 0.0
    shed_mask = np.zeros((n_b, n_h), dtype=float)
    shed_mask[:, 6:8] = 1.0

    model = Model("hourly_54bus_exact_economic")
    log_handle = log_path.open("w", encoding="utf-8", buffering=1)
    model.setLogHandler(log_handle)
    model.setSolverParam("mioMaxTime", float(time_limit))
    model.setSolverParam("optimizerMaxTime", float(time_limit))
    model.setSolverParam("numThreads", int(threads))
    model.setSolverParam("mioSeed", 240924)
    model.setSolverParam("mioHeuristicLevel", 100)
    model.setSolverParam("mioFeaspumpLevel", 2)
    model.setSolverParam("mioRensMaxNodes", 100)

    exact_integer_names = set() if exact_integer_names is None else set(exact_integer_names)

    def binary_domain(name):
        if relax_integrality and name not in exact_integer_names:
            return Domain.inRange(0.0, 1.0)
        return Domain.binary()

    def bounded_integer_domain(name, lower, upper):
        base = Domain.inRange(float(lower), float(upper))
        return base if relax_integrality and name not in exact_integer_names else Domain.integral(base)

    # Planning variables.
    uL = model.variable("uL", n_e, binary_domain("uL"))
    uNS = model.variable("uNS", [2, 2], binary_domain("uNS"))
    uEX = model.variable("uEX", [2, 2], binary_domain("uEX"))
    zESS = model.variable(
        "zESS",
        7,
        bounded_integer_domain(
            "zESS",
            planning["ess_units_per_bus_min"], planning["ess_units_per_bus_max"]
        ),
    )
    zPVG = model.variable(
        "zPVG",
        7,
        bounded_integer_domain(
            "zPVG",
            planning["pv_units_per_bus_min"], planning["pv_units_per_bus_max"]
        ),
    )
    zWTG = model.variable(
        "zWTG",
        6,
        bounded_integer_domain(
            "zWTG",
            planning["wind_units_per_bus_min"], planning["wind_units_per_bus_max"]
        ),
    )

    # Scenario and hourly variables.
    FF = model.variable("FF", [n_e, n_s], Domain.unbounded())
    FS = model.variable("FS", [n_b, n_s], Domain.greaterThan(0.0))
    P = model.variable("P", [n_e, n_h], Domain.unbounded())
    Q = model.variable("Q", [n_e, n_h], Domain.unbounded())
    ell = model.variable("l", [n_e, n_h], Domain.greaterThan(0.0))
    v = model.variable("v", [n_b, n_h], Domain.unbounded())
    P_s = model.variable("P_s", [4, n_h], Domain.greaterThan(0.0))
    Q_s = model.variable("Q_s", [4, n_h], Domain.unbounded())
    kt = model.variable("kt", [4, n_h], bounded_integer_domain("kt", -8.0, 8.0))
    P_pv = model.variable("P_pv", [7, n_h], Domain.unbounded())
    Q_pv = model.variable("Q_pv", [7, n_h], Domain.unbounded())
    P_wt = model.variable("P_wt", [6, n_h], Domain.unbounded())
    Q_wt = model.variable("Q_wt", [6, n_h], Domain.unbounded())
    E_ess = model.variable("E_ess", [7, n_h], Domain.unbounded())
    P_ess = model.variable("P_ess", [7, n_h], Domain.unbounded())
    P_ess_cha = model.variable("P_ess_cha", [7, n_h], Domain.greaterThan(0.0))
    P_ess_dis = model.variable("P_ess_dis", [7, n_h], Domain.greaterThan(0.0))
    u_ess = model.variable("u_ess", [7, n_h], binary_domain("u_ess"))
    P_sop_1 = model.variable("P_sop_1", [4, n_h], Domain.unbounded())
    P_sop_2 = model.variable("P_sop_2", [4, n_h], Domain.unbounded())
    Q_sop_1 = model.variable("Q_sop_1", [4, n_h], Domain.greaterThan(0.0))
    Q_sop_2 = model.variable("Q_sop_2", [4, n_h], Domain.greaterThan(0.0))
    Pl_sop_1 = model.variable("Pl_sop_1", [4, n_h], Domain.greaterThan(0.0))
    Pl_sop_2 = model.variable("Pl_sop_2", [4, n_h], Domain.greaterThan(0.0))
    P_ls = model.variable("P_ls", [n_b, n_h], Domain.greaterThan(0.0))

    variables = {
        "uL": (uL, (63, 1)),
        "uNS": (uNS, (2, 2)),
        "uEX": (uEX, (2, 2)),
        "zESS": (zESS, (7, 1)),
        "zPVG": (zPVG, (7, 1)),
        "zWTG": (zWTG, (6, 1)),
        "u_ess": (u_ess, (7, 96)),
        "kt": (kt, (4, 96)),
        "FF": (FF, (63, 4)),
        "FS": (FS, (54, 4)),
        "P": (P, (63, 96)),
        "Q": (Q, (63, 96)),
        "l": (ell, (63, 96)),
        "v": (v, (54, 96)),
        "P_s": (P_s, (4, 96)),
        "Q_s": (Q_s, (4, 96)),
        "P_pv": (P_pv, (7, 96)),
        "Q_pv": (Q_pv, (7, 96)),
        "P_wt": (P_wt, (6, 96)),
        "Q_wt": (Q_wt, (6, 96)),
        "E_ess": (E_ess, (7, 96)),
        "P_ess": (P_ess, (7, 96)),
        "P_ess_cha": (P_ess_cha, (7, 96)),
        "P_ess_dis": (P_ess_dis, (7, 96)),
        "P_sop_1": (P_sop_1, (4, 96)),
        "P_sop_2": (P_sop_2, (4, 96)),
        "Q_sop_1": (Q_sop_1, (4, 96)),
        "Q_sop_2": (Q_sop_2, (4, 96)),
        "Pl_sop_1": (Pl_sop_1, (4, 96)),
        "Pl_sop_2": (Pl_sop_2, (4, 96)),
        "P_ls": (P_ls, (54, 96)),
    }

    if integer_fix:
        internal_shapes = {
            "uL": (63,),
            "uNS": (2, 2),
            "uEX": (2, 2),
            "zESS": (7,),
            "zPVG": (7,),
            "zWTG": (6,),
            "u_ess": (7, 96),
            "kt": (4, 96),
        }
        for name, expected_shape in internal_shapes.items():
            if name not in integer_fix:
                continue
            variable = variables[name][0]
            values = np.asarray(integer_fix[name], dtype=float).reshape(expected_shape)
            for index in np.ndindex(expected_shape):
                component = variable.index(index[0]) if len(index) == 1 else variable.index(*index)
                eq(model, component, float(values[index]))

    # Planning constraints.
    for e in existing_lines:
        eq(model, uL.index(e), 1.0)
    le(model, Expr.add(Expr.sum(uL), Expr.sum(uNS)), 52.0)
    le(model, Expr.sum(zESS), float(planning["ess_units_total_max"]))
    le(model, Expr.sum(zPVG), float(planning["pv_units_total_max"]))
    for i in range(2):
        le(model, Expr.add(uEX.index(i, 0), uEX.index(i, 1)), 1.0)
        le(model, Expr.add(uNS.index(i, 0), uNS.index(i, 1)), 1.0)

    # Topology and connectivity, including repeated scenario cardinality.
    new_count = Expr.sum(uNS)
    for s in range(n_s):
        eq(model, Expr.add(Expr.sum(uL), new_count), 52.0)
        for e in range(n_e):
            le(model, Expr.sub(FF.index(e, s), Expr.mul(54.0, uL.index(e))), 0.0)
            le(model, Expr.sub(Expr.neg(FF.index(e, s)), Expr.mul(54.0, uL.index(e))), 0.0)
        for b in range(n_b):
            balance_terms = [FS.index(b, s), -1.0]
            balance_terms.extend(FF.index(e, s) for e in incoming[b])
            balance_terms.extend(Expr.neg(FF.index(e, s)) for e in outgoing[b])
            eq(model, add_terms(balance_terms), 0.0)
            if b < 50:
                eq(model, FS.index(b, s), 0.0)
            elif b < 52:
                gate = Expr.add(uEX.index(b - 50, 0), uEX.index(b - 50, 1))
                le(model, Expr.sub(FS.index(b, s), Expr.mul(54.0, gate)), 0.0)
            else:
                gate = Expr.add(uNS.index(b - 52, 0), uNS.index(b - 52, 1))
                le(model, Expr.sub(FS.index(b, s), Expr.mul(54.0, gate)), 0.0)

    # Station capacities, voltage and OLTC constraints.
    station_cap = [
        add_terms([1.67, Expr.mul(1.33, uEX.index(0, 0)), Expr.mul(1.67, uEX.index(0, 1))]),
        add_terms([1.67, Expr.mul(1.33, uEX.index(1, 0)), Expr.mul(1.67, uEX.index(1, 1))]),
        add_terms([Expr.mul(1.67, uNS.index(0, 0)), Expr.mul(2.22, uNS.index(0, 1))]),
        add_terms([Expr.mul(1.67, uNS.index(1, 0)), Expr.mul(2.22, uNS.index(1, 1))]),
    ]
    for r in range(4):
        for h in range(n_h):
            derate = 0.30 if r < 2 and h in (6, 7) else 1.0
            rhs = Expr.mul(0.5 * derate, station_cap[r])
            model.constraint(
                Expr.vstack(rhs, P_s.index(r, h), Q_s.index(r, h)), Domain.inQCone()
            )
    for b in range(50):
        for h in range(n_h):
            model.constraint(v.index(b, h), Domain.inRange(v_min_sq, v_max_sq))
    for r in range(2):
        for h in range(n_h):
            eq(
                model,
                Expr.sub(v.index(50 + r, h), Expr.mul(0.025, kt.index(r, h))),
                1.01,
            )
    for i in range(2):
        new_on = Expr.add(uNS.index(i, 0), uNS.index(i, 1))
        b = 52 + i
        r = 2 + i
        for h in range(n_h):
            tap_voltage = Expr.add(1.01, Expr.mul(0.025, kt.index(r, h)))
            # tap_voltage - M*(1-new_on) <= v <= tap_voltage + M*(1-new_on)
            le(
                model,
                add_terms([tap_voltage, -big_m, Expr.mul(big_m, new_on), Expr.neg(v.index(b, h))]),
                0.0,
            )
            le(
                model,
                add_terms([v.index(b, h), Expr.neg(tap_voltage), -big_m, Expr.mul(big_m, new_on)]),
                0.0,
            )
            le(model, add_terms([v_min_sq, Expr.mul(-big_m, new_on), Expr.neg(v.index(b, h))]), 0.0)
            le(model, add_terms([v.index(b, h), -v_max_sq, Expr.mul(-big_m, new_on)]), 0.0)
            le(model, add_terms([kt.index(r, h), Expr.mul(-8.0, new_on)]), 0.0)
            le(model, add_terms([Expr.neg(kt.index(r, h)), Expr.mul(-8.0, new_on)]), 0.0)

    # Branch bounds, exact voltage drops, and SOC branch-flow relaxation.
    for e in range(n_e):
        for h in range(n_h):
            y_coeff = mask[e, h]
            y_eff = Expr.mul(y_coeff, uL.index(e))
            le(model, Expr.sub(ell.index(e, h), Expr.mul(i_sq_max, y_eff)), 0.0)
            le(model, Expr.sub(P.index(e, h), Expr.mul(s_line_max, y_eff)), 0.0)
            le(model, Expr.sub(Expr.neg(P.index(e, h)), Expr.mul(s_line_max, y_eff)), 0.0)
            le(model, Expr.sub(Q.index(e, h), Expr.mul(s_line_max, y_eff)), 0.0)
            le(model, Expr.sub(Expr.neg(Q.index(e, h)), Expr.mul(s_line_max, y_eff)), 0.0)
            model.constraint(
                Expr.vstack(
                    Expr.vstack(
                        Expr.mul(0.5, v.index(fr[e], h)),
                        ell.index(e, h),
                        P.index(e, h),
                    ),
                    Q.index(e, h),
                ),
                Domain.inRotatedQCone(),
            )
            drop = add_terms(
                [
                    v.index(fr[e], h),
                    Expr.neg(v.index(to[e], h)),
                    Expr.mul(-2.0 * resistance[e], P.index(e, h)),
                    Expr.mul(-2.0 * reactance[e], Q.index(e, h)),
                    Expr.mul(resistance[e] ** 2 + reactance[e] ** 2, ell.index(e, h)),
                ]
            )
            le(model, add_terms([drop, Expr.mul(big_m, y_eff)]), big_m)
            le(model, add_terms([Expr.neg(drop), Expr.mul(big_m, y_eff)]), big_m)

    # Wind and PV generation.
    for i in range(6):
        for h in range(n_h):
            avail = Expr.mul(sdg_rated * wt_fc[h], zWTG.index(i))
            ge(model, Expr.sub(P_wt.index(i, h), Expr.mul(0.3, avail)), 0.0)
            le(model, Expr.sub(P_wt.index(i, h), avail), 0.0)
            model.constraint(
                Expr.vstack(Expr.mul(sdg_rated, zWTG.index(i)), P_wt.index(i, h), Q_wt.index(i, h)),
                Domain.inQCone(),
            )
            le(model, Expr.sub(Q_wt.index(i, h), Expr.mul(qfac, P_wt.index(i, h))), 0.0)
            le(model, Expr.sub(Expr.neg(Q_wt.index(i, h)), Expr.mul(qfac, P_wt.index(i, h))), 0.0)
    for i in range(7):
        for h in range(n_h):
            avail = Expr.mul(sdg_rated * pv_fc[h], zPVG.index(i))
            ge(model, Expr.sub(P_pv.index(i, h), Expr.mul(0.3, avail)), 0.0)
            le(model, Expr.sub(P_pv.index(i, h), avail), 0.0)
            model.constraint(
                Expr.vstack(Expr.mul(sdg_rated, zPVG.index(i)), P_pv.index(i, h), Q_pv.index(i, h)),
                Domain.inQCone(),
            )
            le(model, Expr.sub(Q_pv.index(i, h), Expr.mul(qfac, P_pv.index(i, h))), 0.0)
            le(model, Expr.sub(Expr.neg(Q_pv.index(i, h)), Expr.mul(qfac, P_pv.index(i, h))), 0.0)

    # Storage power, mode, energy bounds, and authoritative wrap-around chronology.
    for i in range(7):
        cap = Expr.mul(ess_unit_pu[i], zESS.index(i))
        pmax = Expr.mul(float(operation["ess_power_to_energy_ratio"]), cap)
        e0 = Expr.mul(float(operation["ess_soc_initial"]), cap)
        for h in range(n_h):
            eq(
                model,
                add_terms(
                    [P_ess.index(i, h), Expr.neg(P_ess_dis.index(i, h)), P_ess_cha.index(i, h)]
                ),
                0.0,
            )
            le(model, Expr.sub(P_ess_cha.index(i, h), pmax), 0.0)
            le(model, Expr.sub(P_ess_dis.index(i, h), pmax), 0.0)
            le(model, Expr.sub(P_ess_cha.index(i, h), Expr.mul(big_m, u_ess.index(i, h))), 0.0)
            le(
                model,
                add_terms([P_ess_dis.index(i, h), -big_m, Expr.mul(big_m, u_ess.index(i, h))]),
                0.0,
            )
            ge(model, Expr.sub(E_ess.index(i, h), Expr.mul(float(operation["ess_soc_min"]), cap)), 0.0)
            le(model, Expr.sub(E_ess.index(i, h), cap), 0.0)
        for s in range(n_s):
            base = 24 * s
            for t in range(23):
                eq(
                    model,
                    add_terms(
                        [
                            E_ess.index(i, base + t + 1),
                            Expr.neg(E_ess.index(i, base + t)),
                            Expr.mul(-eta, P_ess_cha.index(i, base + t)),
                            Expr.mul(1.0 / eta, P_ess_dis.index(i, base + t)),
                        ]
                    ),
                    0.0,
                )
            eq(
                model,
                add_terms(
                    [
                        E_ess.index(i, base),
                        Expr.neg(e0),
                        Expr.mul(-eta, P_ess_cha.index(i, base + 23)),
                        Expr.mul(1.0 / eta, P_ess_dis.index(i, base + 23)),
                    ]
                ),
                0.0,
            )
            eq(model, Expr.sub(E_ess.index(i, base + 23), e0), 0.0)

    # SOP constraints; capacities and investment are fixed regardless of line build.
    z_sop_1 = np.asarray(planning["sop_units_endpoint_1"], dtype=float)
    z_sop_2 = np.asarray(planning["sop_units_endpoint_2"], dtype=float)
    sop_unit_pu = float(planning["sop_unit_mva"]) / s_base
    loss_factor = float(operation["sop_loss_factor"])
    q_fraction = float(operation["sop_q_upper_fraction"])
    for m in range(4):
        s1 = z_sop_1[m] * sop_unit_pu
        s2 = z_sop_2[m] * sop_unit_pu
        for h in range(n_h):
            model.constraint(P_sop_1.index(m, h), Domain.inRange(-s1, s1))
            model.constraint(P_sop_2.index(m, h), Domain.inRange(-s2, s2))
            le(model, Q_sop_1.index(m, h), q_fraction * s1)
            le(model, Q_sop_2.index(m, h), q_fraction * s2)
            le(model, Pl_sop_1.index(m, h), loss_factor * s1)
            le(model, Pl_sop_2.index(m, h), loss_factor * s2)
            model.constraint(
                Expr.vstack(Expr.mul(1.0 / loss_factor, Pl_sop_1.index(m, h)), P_sop_1.index(m, h), Q_sop_1.index(m, h)),
                Domain.inQCone(),
            )
            model.constraint(
                Expr.vstack(s1, P_sop_1.index(m, h), Q_sop_1.index(m, h)), Domain.inQCone()
            )
            model.constraint(
                Expr.vstack(Expr.mul(1.0 / loss_factor, Pl_sop_2.index(m, h)), P_sop_2.index(m, h), Q_sop_2.index(m, h)),
                Domain.inQCone(),
            )
            model.constraint(
                Expr.vstack(s2, P_sop_2.index(m, h), Q_sop_2.index(m, h)), Domain.inQCone()
            )
            eq(
                model,
                add_terms(
                    [
                        P_sop_1.index(m, h),
                        P_sop_2.index(m, h),
                        Pl_sop_1.index(m, h),
                        Pl_sop_2.index(m, h),
                    ]
                ),
                0.0,
            )

    # Load shedding and nodal active/reactive balances.
    for b in range(n_b):
        for h in range(n_h):
            upper = pload[b, h] * shed_mask[b, h]
            if b == 49 and h in (6, 7):
                upper = min(upper, 0.10 * pload[b, h])
            le(model, P_ls.index(b, h), upper)

            p_terms = [-float(pload[b, h]), P_ls.index(b, h)]
            q_terms = [-float(qload[b, h]), Expr.mul(tan_phi, P_ls.index(b, h))]
            for e in outgoing[b]:
                p_terms.append(Expr.neg(P.index(e, h)))
                q_terms.append(Expr.neg(Q.index(e, h)))
            for e in incoming[b]:
                p_terms.extend([P.index(e, h), Expr.mul(resistance[e], ell.index(e, h))])
                q_terms.extend([Q.index(e, h), Expr.mul(reactance[e], ell.index(e, h))])
            if b in station_buses:
                r = station_buses.index(b)
                p_terms.append(P_s.index(r, h))
                q_terms.append(Q_s.index(r, h))
            if b in wt_buses:
                i = wt_buses.index(b)
                p_terms.append(P_wt.index(i, h))
                q_terms.append(Q_wt.index(i, h))
            if b in pv_buses:
                i = pv_buses.index(b)
                p_terms.append(P_pv.index(i, h))
                q_terms.append(Q_pv.index(i, h))
            if b in ess_buses:
                i = ess_buses.index(b)
                p_terms.append(P_ess.index(i, h))
            for m, e in enumerate(sop_lines):
                if b == fr[e]:
                    p_terms.append(P_sop_1.index(m, h))
                    q_terms.append(Q_sop_1.index(m, h))
                if b == to[e]:
                    p_terms.append(P_sop_2.index(m, h))
                    q_terms.append(Q_sop_2.index(m, h))
            eq(model, add_terms(p_terms), 0.0)
            eq(model, add_terms(q_terms), 0.0)

    # Exact annualized economic objective in dollars/year.
    years = int(planning["years"])
    discount = float(planning["discount_rate"])
    c_dn = discount * (1.0 + discount) ** years / ((1.0 + discount) ** years - 1.0)
    ex_cost = np.asarray(planning["existing_station_option_cost_usd"], dtype=float)
    ns_cost = np.asarray(planning["candidate_station_option_cost_usd"], dtype=float)
    line_investment = float(planning["line_build_cost_usd_per_km"]) * lengths.copy()
    line_investment[list(existing_lines)] = 0.0
    sop_fixed_investment = (
        float(economics["sop_investment_usd_per_mva"])
        * float(planning["sop_unit_mva"])
        * float(np.sum(z_sop_1 + z_sop_2))
    )
    c_inv = add_terms(
        [
        Expr.dot(ex_cost.reshape(-1), uEX.reshape(4)),
        Expr.dot(ns_cost.reshape(-1), uNS.reshape(4)),
        Expr.dot(line_investment, uL),
        Expr.mul(float(economics["wind_investment_usd_per_mw"]) * float(planning["dg_unit_mw"]), Expr.sum(zWTG))
        ,
        Expr.mul(float(economics["pv_investment_usd_per_mw"]) * float(planning["dg_unit_mw"]), Expr.sum(zPVG))
        ,
        sop_fixed_investment,
        Expr.dot(float(economics["ess_investment_usd_per_mwh"]) * ess_unit_mwh, zESS),
        ]
    )

    maintenance_terms = [
        Expr.mul(
            float(planning["dg_unit_mw"])
            * float(economics["pv_maintenance_usd_per_mwh"])
            * float(economics["pv_capacity_factor"])
            * 8760.0,
            Expr.sum(zPVG),
        ),
        Expr.mul(
            float(planning["dg_unit_mw"])
            * float(economics["wind_maintenance_usd_per_mwh"])
            * float(economics["wind_capacity_factor"])
            * 8760.0,
            Expr.sum(zWTG),
        ),
        float(np.sum(z_sop_1 + z_sop_2))
        * float(planning["sop_unit_mva"])
        * float(economics["sop_maintenance_fraction"])
        * float(economics["sop_investment_usd_per_mva"]),
    ]
    # 0.25*sum over four scenarios is exactly one for scenario-invariant uL.
    maintenance_terms.append(
        Expr.dot(float(planning["line_maintenance_usd_per_km_year"]) * lengths, uL)
    )
    c_ma = add_terms(maintenance_terms)

    price = np.asarray(operation["purchase_price_usd_per_mwh"], dtype=float)
    annual_factor = float(economics["scenario_day_annualization_factor"])
    curtail_coeff = s_base * float(economics["curtailment_cost_usd_per_mwh"])
    shed_coeff = s_base * float(economics["load_shed_cost_usd_per_mwh"])
    price_96 = np.tile(price, 4)
    c_ope = add_terms(
        [
            Expr.dot(np.tile(s_base * price_96, 4), P_s.reshape(4 * n_h)),
            Expr.mul(curtail_coeff * sdg_rated * float(np.sum(wt_fc)), Expr.sum(zWTG)),
            Expr.mul(-curtail_coeff, Expr.sum(P_wt)),
            Expr.mul(curtail_coeff * sdg_rated * float(np.sum(pv_fc)), Expr.sum(zPVG)),
            Expr.mul(-curtail_coeff, Expr.sum(P_pv)),
            Expr.mul(shed_coeff, Expr.sum(P_ls)),
        ]
    )
    objective = add_terms([Expr.mul(c_dn, c_inv), Expr.mul(annual_factor, c_ope), c_ma])
    model.objective("annual_cost", ObjectiveSense.Minimize, objective)

    metadata = {
        "c_dn": c_dn,
        "sop_fixed_investment": sop_fixed_investment,
        "model_dimensions": {"buses": n_b, "branches": n_e, "hours": n_h, "scenarios": n_s},
    }
    return model, variables, log_handle, metadata


def solver_info(model, key, default=None):
    try:
        return float(model.getSolverDoubleInfo(key))
    except Exception:
        return default


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--time-limit", type=float, default=870.0)
    parser.add_argument("--threads", type=int, default=4)
    args = parser.parse_args()
    if args.time_limit <= 0 or args.time_limit > 900:
        raise ValueError("time limit must be in (0, 900]")
    if args.threads < 1 or args.threads > 4:
        raise ValueError("threads must be in [1, 4]")

    result_dir = args.output_root / "results"
    log_dir = args.output_root / "logs"
    result_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)
    spec, branch, bus, forecast = read_inputs(args.input_root)

    build_start = time.monotonic()
    model, variables, log_handle, metadata = build_model(
        spec, branch, bus, forecast, log_dir / "mosek.log", args.time_limit, args.threads
    )
    build_seconds = time.monotonic() - build_start
    solve_start = time.monotonic()
    exception_text = None
    try:
        model.solve()
    except Exception as exc:
        exception_text = repr(exc)
    wall_solve_seconds = time.monotonic() - solve_start

    primal_status = str(model.getPrimalSolutionStatus())
    problem_status = str(model.getProblemStatus())
    has_solution = model.getPrimalSolutionStatus() in (
        SolutionStatus.Optimal,
        SolutionStatus.Feasible,
    )
    objective = None
    best_bound = solver_info(model, "mioObjBound")
    relative_gap = solver_info(model, "mioObjRelGap")
    optimizer_time = solver_info(model, "optimizerTime", wall_solve_seconds)
    if has_solution:
        model.acceptedSolutionStatus(AccSolutionStatus.Feasible)
        objective = float(model.primalObjValue())
        arrays = {}
        for name, (variable, shape) in variables.items():
            values = np.asarray(variable.level(), dtype=np.float64).reshape(shape)
            arrays[name] = values
        np.savez_compressed(result_dir / "solution.npz", **arrays)

    if has_solution and model.getPrimalSolutionStatus() == SolutionStatus.Optimal:
        status = "optimal"
    elif has_solution:
        status = "feasible"
    else:
        status = "failed" if exception_text else "no_incumbent"

    summary = {
        "status": status,
        "solver": "MOSEK 11.2 Fusion MISOCP",
        "optimizer_time_seconds": optimizer_time,
        "threads": args.threads,
        "objective": objective,
        "best_bound": best_bound,
        "relative_gap": relative_gap,
        "solution_path": "candidate/results/solution.npz",
        "used_warm_start": False,
        "used_external_warm_start": False,
        "internal_mip_start_source": None,
        "planning_was_fixed": False,
        "optimizer_time_limit_seconds": args.time_limit,
        "build_time_seconds": build_seconds,
        "solve_wall_time_seconds": wall_solve_seconds,
        "problem_status": problem_status,
        "primal_solution_status": primal_status,
        "exception": exception_text,
        "random_seed": 240924,
        "metadata": metadata,
    }
    with (result_dir / "solver_summary.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2, sort_keys=True)
        handle.write("\n")

    log_handle.close()
    model.dispose()
    print(json.dumps(summary, indent=2, sort_keys=True))
    if not has_solution:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
