#!/usr/bin/env python3
"""Independent replay auditor for the exported candidate.

This module intentionally does not import the model-building implementation.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import numpy as np
import openpyxl


REQUIRED_SHAPES = {
    "uL": (63, 1),
    "uNS": (2, 2),
    "uEX": (2, 2),
    "zESS": (7, 1),
    "zPVG": (7, 1),
    "zWTG": (6, 1),
    "u_ess": (7, 96),
    "kt": (4, 96),
    "FF": (63, 4),
    "FS": (54, 4),
    "P": (63, 96),
    "Q": (63, 96),
    "l": (63, 96),
    "v": (54, 96),
    "P_s": (4, 96),
    "Q_s": (4, 96),
    "P_pv": (7, 96),
    "Q_pv": (7, 96),
    "P_wt": (6, 96),
    "Q_wt": (6, 96),
    "E_ess": (7, 96),
    "P_ess": (7, 96),
    "P_ess_cha": (7, 96),
    "P_ess_dis": (7, 96),
    "P_sop_1": (4, 96),
    "P_sop_2": (4, 96),
    "Q_sop_1": (4, 96),
    "Q_sop_2": (4, 96),
    "Pl_sop_1": (4, 96),
    "Pl_sop_2": (4, 96),
    "P_ls": (54, 96),
}


def max_abs(value):
    arr = np.asarray(value, dtype=float)
    return float(np.max(np.abs(arr))) if arr.size else 0.0


def max_pos(value):
    arr = np.asarray(value, dtype=float)
    return float(max(0.0, np.max(arr))) if arr.size else 0.0


def max_bound_violation(value, lower=None, upper=None):
    result = 0.0
    if lower is not None:
        result = max(result, max_pos(np.asarray(lower) - np.asarray(value)))
    if upper is not None:
        result = max(result, max_pos(np.asarray(value) - np.asarray(upper)))
    return result


def read_inputs(input_root: Path):
    with (input_root / "model_spec.json").open("r", encoding="utf-8") as handle:
        spec = json.load(handle)
    workbook = openpyxl.load_workbook(
        input_root / "line_load.xlsx", read_only=True, data_only=True
    )
    ws1 = workbook["Sheet1"]
    ws3 = workbook["Sheet3"]
    branch = np.asarray(
        [[ws1.cell(row, col).value for col in range(1, 5)] for row in range(2, 65)],
        dtype=np.float64,
    )
    bus = np.asarray(
        [[ws1.cell(row, col).value for col in range(7, 9)] for row in range(2, 56)],
        dtype=np.float64,
    )
    forecast = np.asarray(
        [[ws3.cell(row, col).value for col in range(2, 5)] for row in range(2, 98)],
        dtype=np.float64,
    )
    workbook.close()
    return spec, branch, bus, forecast


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-root", type=Path, required=True)
    parser.add_argument("--candidate-root", type=Path, required=True)
    parser.add_argument("--constraint-tolerance", type=float, default=1.0e-5)
    parser.add_argument("--integrality-tolerance", type=float, default=1.0e-5)
    args = parser.parse_args()

    spec, branch, bus, forecast = read_inputs(args.input_root)
    solution_path = args.candidate_root / "results" / "solution.npz"
    summary_path = args.candidate_root / "results" / "solver_summary.json"
    arrays = {}
    shape_checks = {}
    dtype_checks = {}
    finite_checks = {}
    with np.load(solution_path, allow_pickle=False) as solution:
        for name, shape in REQUIRED_SHAPES.items():
            if name not in solution:
                shape_checks[name] = False
                dtype_checks[name] = False
                finite_checks[name] = False
                continue
            value = solution[name]
            arrays[name] = value
            shape_checks[name] = value.shape == shape
            dtype_checks[name] = value.dtype == np.float64
            finite_checks[name] = bool(np.all(np.isfinite(value)))
    contract_pass = bool(
        len(arrays) == len(REQUIRED_SHAPES)
        and all(shape_checks.values())
        and all(dtype_checks.values())
        and all(finite_checks.values())
    )
    if not contract_pass:
        report = {
            "audit_passed": False,
            "reason": "array contract failure",
            "array_contract": {
                "shape_checks": shape_checks,
                "float64_checks": dtype_checks,
                "finite_checks": finite_checks,
            },
        }
        with (args.candidate_root / "audit_report.json").open("w", encoding="utf-8") as handle:
            json.dump(report, handle, indent=2, sort_keys=True)
            handle.write("\n")
        raise SystemExit(2)

    with summary_path.open("r", encoding="utf-8") as handle:
        summary = json.load(handle)

    uL = arrays["uL"][:, 0]
    uNS, uEX = arrays["uNS"], arrays["uEX"]
    zESS, zPVG, zWTG = arrays["zESS"][:, 0], arrays["zPVG"][:, 0], arrays["zWTG"][:, 0]
    u_ess, kt = arrays["u_ess"], arrays["kt"]
    FF, FS = arrays["FF"], arrays["FS"]
    P, Q, ell, v = arrays["P"], arrays["Q"], arrays["l"], arrays["v"]
    P_s, Q_s = arrays["P_s"], arrays["Q_s"]
    P_pv, Q_pv, P_wt, Q_wt = arrays["P_pv"], arrays["Q_pv"], arrays["P_wt"], arrays["Q_wt"]
    E_ess, P_ess = arrays["E_ess"], arrays["P_ess"]
    P_ess_cha, P_ess_dis = arrays["P_ess_cha"], arrays["P_ess_dis"]
    P_sop_1, P_sop_2 = arrays["P_sop_1"], arrays["P_sop_2"]
    Q_sop_1, Q_sop_2 = arrays["Q_sop_1"], arrays["Q_sop_2"]
    Pl_sop_1, Pl_sop_2 = arrays["Pl_sop_1"], arrays["Pl_sop_2"]
    P_ls = arrays["P_ls"]

    network = spec["network"]
    planning = spec["planning"]
    operation = spec["operation"]
    economics = spec["economics"]
    n_b, n_e, n_h, n_s = 54, 63, 96, 4
    big_m = float(network["big_m"])
    s_base = float(network["s_base_mva"])
    v_base = float(network["v_base_kv"])
    v_min_sq = float(network["v_min_pu"]) ** 2
    v_max_sq = float(network["v_max_pu"]) ** 2
    i_max = float(network["line_current_limit_ka"]) / (s_base / v_base)
    i_sq_max = i_max**2
    s_line_max = i_max * float(network["v_max_pu"])
    z_base = v_base**2 / s_base
    lengths = branch[:, 3]
    fr = branch[:, 1].astype(int) - 1
    to = branch[:, 2].astype(int) - 1
    resistance = float(network["resistance_ohm_per_km"]) * lengths / z_base
    reactance = float(network["reactance_ohm_per_km"]) * lengths / z_base
    wt_fc, pv_fc, pl_fc = forecast[:, 0], forecast[:, 1], forecast[:, 2]
    pload = bus[:, 1, None] * pl_fc[None, :] / s_base
    tan_phi = math.tan(math.acos(float(network["load_power_factor"])))
    qload = pload * tan_phi
    qfac = math.sqrt(1.0 - float(network["dg_min_power_factor"]) ** 2) / float(network["dg_min_power_factor"])
    sdg_rated = float(planning["dg_unit_mw"]) / s_base
    eta = float(operation["ess_efficiency"])
    ess_unit_mwh = np.full(7, float(operation["ess_unit_mwh_non_idc"]), dtype=float)
    ess_unit_mwh[-1] = float(operation["ess_unit_mwh_idc"])
    ess_unit_pu = ess_unit_mwh / s_base
    existing_lines = np.asarray([int(e) - 1 for e in spec["sets"]["existing_line_rows"]], dtype=int)
    fault_lines = np.asarray([int(e) - 1 for e in spec["sets"]["fault_line_rows"]], dtype=int)
    wt_buses = np.asarray([int(b) - 1 for b in spec["sets"]["wind_buses"]], dtype=int)
    pv_buses = np.asarray([int(b) - 1 for b in spec["sets"]["pv_buses"]], dtype=int)
    ess_buses = np.asarray([int(b) - 1 for b in spec["sets"]["ess_buses"]], dtype=int)
    sop_lines = np.asarray([int(e) - 1 for e in spec["sets"]["sop_line_rows"]], dtype=int)

    A_fr = np.zeros((n_b, n_e), dtype=float)
    A_to = np.zeros((n_b, n_e), dtype=float)
    A_fr[fr, np.arange(n_e)] = 1.0
    A_to[to, np.arange(n_e)] = 1.0
    mask = np.ones((n_e, n_h), dtype=float)
    mask[fault_lines[:, None], np.asarray([6, 7])[None, :]] = 0.0
    y_eff = uL[:, None] * mask
    shed_mask = np.zeros((n_b, n_h), dtype=float)
    shed_mask[:, 6:8] = 1.0

    residuals = {}

    def rec(name, value):
        residuals[name] = float(value)

    # Every variable domain and every integer family.
    rec("domain.uL", max_bound_violation(uL, 0.0, 1.0))
    rec("domain.uNS", max_bound_violation(uNS, 0.0, 1.0))
    rec("domain.uEX", max_bound_violation(uEX, 0.0, 1.0))
    rec("domain.zESS", max_bound_violation(zESS, planning["ess_units_per_bus_min"], planning["ess_units_per_bus_max"]))
    rec("domain.zPVG", max_bound_violation(zPVG, planning["pv_units_per_bus_min"], planning["pv_units_per_bus_max"]))
    rec("domain.zWTG", max_bound_violation(zWTG, planning["wind_units_per_bus_min"], planning["wind_units_per_bus_max"]))
    rec("domain.u_ess", max_bound_violation(u_ess, 0.0, 1.0))
    rec("domain.kt", max_bound_violation(kt, -8.0, 8.0))
    rec("domain.FS_nonnegative", max_pos(-FS))
    rec("domain.l_nonnegative", max_pos(-ell))
    rec("domain.P_s_nonnegative", max_pos(-P_s))
    rec("domain.P_ess_cha_nonnegative", max_pos(-P_ess_cha))
    rec("domain.P_ess_dis_nonnegative", max_pos(-P_ess_dis))
    rec("domain.Q_sop_1_nonnegative", max_pos(-Q_sop_1))
    rec("domain.Q_sop_2_nonnegative", max_pos(-Q_sop_2))
    rec("domain.Pl_sop_1_nonnegative", max_pos(-Pl_sop_1))
    rec("domain.Pl_sop_2_nonnegative", max_pos(-Pl_sop_2))
    rec("domain.P_ls_nonnegative", max_pos(-P_ls))
    integer_families = {
        "uL": uL,
        "uNS": uNS,
        "uEX": uEX,
        "zESS": zESS,
        "zPVG": zPVG,
        "zWTG": zWTG,
        "u_ess": u_ess,
        "kt": kt,
    }
    integrality = {name: max_abs(value - np.rint(value)) for name, value in integer_families.items()}
    max_integrality = max(integrality.values())

    # Planning limits and option selection.
    rec("planning.existing_line_fix", max_abs(uL[existing_lines] - 1.0))
    rec("planning.line_count_upper", max_pos(np.sum(uL) + np.sum(uNS) - 52.0))
    rec("planning.ess_total", max_pos(np.sum(zESS) - float(planning["ess_units_total_max"])))
    rec("planning.pv_total", max_pos(np.sum(zPVG) - float(planning["pv_units_total_max"])))
    rec("planning.existing_option_selection", max_pos(np.sum(uEX, axis=1) - 1.0))
    rec("planning.new_option_selection", max_pos(np.sum(uNS, axis=1) - 1.0))

    # Topology, connectivity, and exact root gates.
    card = np.full(n_s, np.sum(uL) + np.sum(uNS) - 52.0)
    rec("topology.cardinality", max_abs(card))
    rec("topology.FF_bounds", max_pos(np.abs(FF) - 54.0 * uL[:, None]))
    rec("topology.flow_balance", max_abs((A_to - A_fr) @ FF - 1.0 + FS))
    rec("topology.nonroot_FS_zero", max_abs(FS[:50, :]))
    root_caps = np.vstack(
        [
            54.0 * np.sum(uEX[0, :]) * np.ones(n_s),
            54.0 * np.sum(uEX[1, :]) * np.ones(n_s),
            54.0 * np.sum(uNS[0, :]) * np.ones(n_s),
            54.0 * np.sum(uNS[1, :]) * np.ones(n_s),
        ]
    )
    rec("topology.root_FS_gates", max_pos(FS[50:54, :] - root_caps))

    # Station capacities and derating cones.
    station_cap = np.asarray(
        [
            1.67 + 1.33 * uEX[0, 0] + 1.67 * uEX[0, 1],
            1.67 + 1.33 * uEX[1, 0] + 1.67 * uEX[1, 1],
            1.67 * uNS[0, 0] + 2.22 * uNS[0, 1],
            1.67 * uNS[1, 0] + 2.22 * uNS[1, 1],
        ]
    )
    station_rhs = 0.5 * station_cap[:, None] * np.ones((4, n_h))
    station_rhs[:2, 6:8] *= 0.30
    rec("station.apparent_power_cones", max_pos(np.hypot(P_s, Q_s) - station_rhs))

    # Voltage and tap semantics.
    rec("voltage.load_bus_bounds", max_bound_violation(v[:50, :], v_min_sq, v_max_sq))
    rec("voltage.existing_tap_equalities", max_abs(v[50:52, :] - (1.01 + 0.025 * kt[:2, :])))
    new_on = np.sum(uNS, axis=1)
    tap_residual = 0.0
    inactive_voltage_residual = 0.0
    tap_gate_residual = 0.0
    for i in range(2):
        tap_voltage = 1.01 + 0.025 * kt[2 + i, :]
        tap_residual = max(tap_residual, max_pos(np.abs(v[52 + i, :] - tap_voltage) - big_m * (1.0 - new_on[i])))
        inactive_voltage_residual = max(
            inactive_voltage_residual,
            max_pos(v_min_sq - big_m * new_on[i] - v[52 + i, :]),
            max_pos(v[52 + i, :] - v_max_sq - big_m * new_on[i]),
        )
        tap_gate_residual = max(tap_gate_residual, max_pos(np.abs(kt[2 + i, :]) - 8.0 * new_on[i]))
    rec("voltage.new_station_tap_big_m", tap_residual)
    rec("voltage.new_station_inactive_bounds", inactive_voltage_residual)
    rec("voltage.new_station_tap_gate", tap_gate_residual)

    # Line limits, outage zeros, voltage drop, and exact SOC expression.
    rec("branch.l_upper", max_pos(ell - i_sq_max * y_eff))
    rec("branch.P_bounds", max_pos(np.abs(P) - s_line_max * y_eff))
    rec("branch.Q_bounds", max_pos(np.abs(Q) - s_line_max * y_eff))
    rec("branch.fault_zero", max(max_abs(P[fault_lines, 6:8]), max_abs(Q[fault_lines, 6:8]), max_abs(ell[fault_lines, 6:8])))
    drop = (
        v[fr, :]
        - v[to, :]
        - 2.0 * resistance[:, None] * P
        - 2.0 * reactance[:, None] * Q
        + (resistance**2 + reactance**2)[:, None] * ell
    )
    rec("branch.voltage_drop_big_m", max_pos(np.abs(drop) - big_m * (1.0 - y_eff)))
    branch_soc_lhs = np.sqrt((2.0 * P) ** 2 + (2.0 * Q) ** 2 + (v[fr, :] - ell) ** 2)
    rec("branch.SOC", max_pos(branch_soc_lhs - (v[fr, :] + ell)))

    # Renewable availability, installed capacity cones, and actual-P power factor.
    pwt_avail = zWTG[:, None] * wt_fc[None, :] * sdg_rated
    ppv_avail = zPVG[:, None] * pv_fc[None, :] * sdg_rated
    rec("renewable.wind_dispatch_lower", max_pos(0.3 * pwt_avail - P_wt))
    rec("renewable.wind_dispatch_upper", max_pos(P_wt - pwt_avail))
    rec("renewable.wind_capacity_cone", max_pos(np.hypot(P_wt, Q_wt) - zWTG[:, None] * sdg_rated))
    rec("renewable.wind_actual_P_power_factor", max_pos(np.abs(Q_wt) - qfac * P_wt))
    rec("renewable.pv_dispatch_lower", max_pos(0.3 * ppv_avail - P_pv))
    rec("renewable.pv_dispatch_upper", max_pos(P_pv - ppv_avail))
    rec("renewable.pv_capacity_cone", max_pos(np.hypot(P_pv, Q_pv) - zPVG[:, None] * sdg_rated))
    rec("renewable.pv_actual_P_power_factor", max_pos(np.abs(Q_pv) - qfac * P_pv))

    # Storage.
    cap = zESS[:, None] * ess_unit_pu[:, None]
    pmax = float(operation["ess_power_to_energy_ratio"]) * cap
    rec("storage.power_definition", max_abs(P_ess - P_ess_dis + P_ess_cha))
    rec("storage.charge_power_limit", max_pos(P_ess_cha - pmax))
    rec("storage.discharge_power_limit", max_pos(P_ess_dis - pmax))
    rec("storage.charge_mode", max_pos(P_ess_cha - big_m * u_ess))
    rec("storage.discharge_mode", max_pos(P_ess_dis - big_m * (1.0 - u_ess)))
    rec("storage.energy_bounds", max_bound_violation(E_ess, float(operation["ess_soc_min"]) * cap, cap))
    dynamics = []
    wrap = []
    terminal = []
    e0 = float(operation["ess_soc_initial"]) * cap[:, 0]
    for s in range(n_s):
        base = 24 * s
        dynamics.append(
            E_ess[:, base + 1 : base + 24]
            - E_ess[:, base : base + 23]
            - eta * P_ess_cha[:, base : base + 23]
            + P_ess_dis[:, base : base + 23] / eta
        )
        wrap.append(
            E_ess[:, base]
            - e0
            - eta * P_ess_cha[:, base + 23]
            + P_ess_dis[:, base + 23] / eta
        )
        terminal.append(E_ess[:, base + 23] - e0)
    rec("storage.chronology", max_abs(np.concatenate(dynamics, axis=1)))
    rec("storage.wraparound_initial", max_abs(np.column_stack(wrap)))
    rec("storage.scenario_terminal", max_abs(np.column_stack(terminal)))

    # SOP bounds, both cones, fixed capacity semantics, and loss balance.
    z_sop_1 = np.asarray(planning["sop_units_endpoint_1"], dtype=float)
    z_sop_2 = np.asarray(planning["sop_units_endpoint_2"], dtype=float)
    sop_unit_pu = float(planning["sop_unit_mva"]) / s_base
    s1 = z_sop_1[:, None] * sop_unit_pu
    s2 = z_sop_2[:, None] * sop_unit_pu
    loss_factor = float(operation["sop_loss_factor"])
    q_fraction = float(operation["sop_q_upper_fraction"])
    rec("sop.P_bounds", max(max_pos(np.abs(P_sop_1) - s1), max_pos(np.abs(P_sop_2) - s2)))
    rec("sop.Q_upper", max(max_pos(Q_sop_1 - q_fraction * s1), max_pos(Q_sop_2 - q_fraction * s2)))
    rec("sop.loss_upper", max(max_pos(Pl_sop_1 - loss_factor * s1), max_pos(Pl_sop_2 - loss_factor * s2)))
    rec("sop.loss_cones", max(max_pos(np.hypot(P_sop_1, Q_sop_1) - Pl_sop_1 / loss_factor), max_pos(np.hypot(P_sop_2, Q_sop_2) - Pl_sop_2 / loss_factor)))
    rec("sop.capacity_cones", max(max_pos(np.hypot(P_sop_1, Q_sop_1) - s1), max_pos(np.hypot(P_sop_2, Q_sop_2) - s2)))
    rec("sop.active_loss_balance", max_abs(P_sop_1 + P_sop_2 + Pl_sop_1 + Pl_sop_2))

    # Load shedding and nodal balances.
    shed_upper = pload * shed_mask
    shed_upper[49, 6:8] = 0.10 * pload[49, 6:8]
    rec("shedding.mask_and_load_upper", max_pos(P_ls - shed_upper))
    rec("shedding.IDC_service", max_pos(P_ls[49, 6:8] - 0.10 * pload[49, 6:8]))
    p_balance = -A_fr @ P + A_to @ P + A_to @ (resistance[:, None] * ell) - (pload - P_ls)
    q_balance = -A_fr @ Q + A_to @ Q + A_to @ (reactance[:, None] * ell) - (qload - tan_phi * P_ls)
    p_balance[50:54, :] += P_s
    q_balance[50:54, :] += Q_s
    p_balance[wt_buses, :] += P_wt
    q_balance[wt_buses, :] += Q_wt
    p_balance[pv_buses, :] += P_pv
    q_balance[pv_buses, :] += Q_pv
    p_balance[ess_buses, :] += P_ess
    for m, e in enumerate(sop_lines):
        p_balance[fr[e], :] += P_sop_1[m, :]
        q_balance[fr[e], :] += Q_sop_1[m, :]
        p_balance[to[e], :] += P_sop_2[m, :]
        q_balance[to[e], :] += Q_sop_2[m, :]
    rec("nodal.active_balance", max_abs(p_balance))
    rec("nodal.reactive_balance", max_abs(q_balance))

    # Independently reconstructed objective components.
    years = int(planning["years"])
    discount = float(planning["discount_rate"])
    c_dn = discount * (1.0 + discount) ** years / ((1.0 + discount) ** years - 1.0)
    ex_cost = np.asarray(planning["existing_station_option_cost_usd"], dtype=float)
    ns_cost = np.asarray(planning["candidate_station_option_cost_usd"], dtype=float)
    line_invest_coeff = float(planning["line_build_cost_usd_per_km"]) * lengths
    line_invest_coeff[existing_lines] = 0.0
    investment = {
        "existing_station_options": float(np.sum(ex_cost * uEX)),
        "candidate_station_options": float(np.sum(ns_cost * uNS)),
        "lines": float(np.dot(line_invest_coeff, uL)),
        "wind": float(economics["wind_investment_usd_per_mw"]) * float(planning["dg_unit_mw"]) * float(np.sum(zWTG)),
        "pv": float(economics["pv_investment_usd_per_mw"]) * float(planning["dg_unit_mw"]) * float(np.sum(zPVG)),
        "sop_fixed": float(economics["sop_investment_usd_per_mva"]) * float(planning["sop_unit_mva"]) * float(np.sum(z_sop_1 + z_sop_2)),
        "storage": float(economics["ess_investment_usd_per_mwh"]) * float(np.dot(ess_unit_mwh, zESS)),
    }
    c_inv = float(sum(investment.values()))
    f_inv = c_dn * c_inv
    maintenance = {
        "pv": float(np.sum(zPVG)) * float(planning["dg_unit_mw"]) * float(economics["pv_maintenance_usd_per_mwh"]) * float(economics["pv_capacity_factor"]) * 8760.0,
        "wind": float(np.sum(zWTG)) * float(planning["dg_unit_mw"]) * float(economics["wind_maintenance_usd_per_mwh"]) * float(economics["wind_capacity_factor"]) * 8760.0,
        "sop_fixed": float(np.sum(z_sop_1 + z_sop_2)) * float(planning["sop_unit_mva"]) * float(economics["sop_maintenance_fraction"]) * float(economics["sop_investment_usd_per_mva"]),
        "lines": float(np.dot(uL, float(planning["line_maintenance_usd_per_km_year"]) * lengths)),
    }
    c_ma = float(sum(maintenance.values()))
    price = np.asarray(operation["purchase_price_usd_per_mwh"], dtype=float)
    operating = {
        "purchases": s_base * float(np.sum(P_s * np.tile(price, 4)[None, :])),
        "wind_curtailment": s_base * float(economics["curtailment_cost_usd_per_mwh"]) * float(np.sum(pwt_avail - P_wt)),
        "pv_curtailment": s_base * float(economics["curtailment_cost_usd_per_mwh"]) * float(np.sum(ppv_avail - P_pv)),
        "load_shedding": s_base * float(economics["load_shed_cost_usd_per_mwh"]) * float(np.sum(P_ls)),
    }
    c_ope = float(sum(operating.values()))
    annual_factor = float(economics["scenario_day_annualization_factor"])
    f_ope = annual_factor * c_ope + c_ma
    objective = f_inv + f_ope
    summary_objective = summary.get("objective")
    objective_difference = None if summary_objective is None else abs(objective - float(summary_objective))

    max_constraint_residual = max(residuals.values())
    objective_tolerance = max(0.05, 1.0e-9 * abs(objective))
    objective_pass = objective_difference is not None and objective_difference <= objective_tolerance
    audit_passed = bool(
        contract_pass
        and max_constraint_residual <= args.constraint_tolerance
        and max_integrality <= args.integrality_tolerance
        and objective_pass
        and summary.get("planning_was_fixed") is False
        and summary.get("used_warm_start") is False
        and summary.get("used_external_warm_start") is False
        and int(summary.get("threads", 999)) <= 4
        and float(summary.get("optimizer_time_limit_seconds", 9999.0)) <= 900.0
    )
    report = {
        "audit_passed": audit_passed,
        "thresholds": {
            "constraint": args.constraint_tolerance,
            "integrality": args.integrality_tolerance,
            "objective_absolute": objective_tolerance,
        },
        "array_contract": {
            "passed": contract_pass,
            "shape_checks": shape_checks,
            "float64_checks": dtype_checks,
            "finite_checks": finite_checks,
        },
        "integrality_distance_by_family": integrality,
        "maximum_integrality_distance": max_integrality,
        "constraint_residuals": residuals,
        "maximum_constraint_residual": max_constraint_residual,
        "objective_components": {
            "capital_recovery_factor": c_dn,
            "investment_components": investment,
            "C_inv": c_inv,
            "F_inv": f_inv,
            "maintenance_components": maintenance,
            "C_ma": c_ma,
            "operating_components_96h": operating,
            "C_ope": c_ope,
            "annualization_factor": annual_factor,
            "F_ope": f_ope,
            "total": objective,
            "solver_reported_total": summary_objective,
            "absolute_difference": objective_difference,
        },
        "eligibility_checks": {
            "planning_was_free": summary.get("planning_was_fixed") is False,
            "no_warm_start": summary.get("used_warm_start") is False,
            "no_external_warm_start": summary.get("used_external_warm_start") is False,
            "threads_at_most_four": int(summary.get("threads", 999)) <= 4,
            "time_limit_at_most_900": float(summary.get("optimizer_time_limit_seconds", 9999.0)) <= 900.0,
        },
    }
    with (args.candidate_root / "audit_report.json").open("w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2, sort_keys=True)
        handle.write("\n")
    print(json.dumps(report, indent=2, sort_keys=True))
    if not audit_passed:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
